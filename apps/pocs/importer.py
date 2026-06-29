"""
POC importer for the company M365 PoC exports (.xlsx or .csv).

Columns are matched to POC fields by **synonym** (so the older "PoC Follow-up"
export and the newer "PoC List" export both work), dates are parsed from several
formats (incl. ``DD/MM/YYYY``), HTML descriptions are flattened to text, and rows
are upserted by the source ``ID`` (``external_id``) so re-importing updates
rather than duplicates. New POCs inherit the global phase blueprint.

Columns without a matching POC field are ignored. People columns are stored as
text (the export has names, not accounts).
"""

import csv
import html
import io
import json
import re
from datetime import date as _date_cls
from datetime import datetime

import openpyxl
from django.db import transaction
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime

from .models import POC
from .services import apply_phase_templates

# External status text → app status (matched case-insensitively).
_STATUS_MAP = {
    "in execution": POC.Status.ACTIVE,
    "ongoing": POC.Status.ACTIVE,
    "active": POC.Status.ACTIVE,
    "closed": POC.Status.COMPLETED,
    "completed": POC.Status.COMPLETED,
    "cancelled": POC.Status.ARCHIVED,
    "identified": POC.Status.DRAFT,
    "in tendering": POC.Status.DRAFT,
    "on hold": POC.Status.DRAFT,
}

# POC field ← accepted header names (lowercased). First match wins.
FIELD_SYNONYMS = {
    "external_id": {"id"},
    "name": {"title", "poc name", "name"},
    "description": {"pilot description", "poc description", "description"},
    "status": {"status", "poc status"},
    "l2_wbs": {"l2 wbs"},
    "initiative": {"initiative"},
    "customer_segment": {"customer segment", "served segments", "served segment", "segment"},
    "customer": {"customer"},
    "leading_organization": {"leading organization", "leading org"},
    "tendering_start": {"tendering start"},
    "tendering_finish": {"tendering finish"},
    "execution_start": {"execution start", "poc start", "start"},
    "execution_finish": {"execution finish", "poc finish", "finish"},
    "bfo_no": {"bfo no"},
    "pilot_requestor": {"pilot requestor", "poc requestor"},
    "opportunity_leader": {"opportunity leader"},
    "ecostruxure_lead": {"ecostruxure lead", "ecostruxure solution owner"},
    "pilot_tender_leader": {"pilot tender leader"},
    "pilot_tender_tl": {"pilot tender tl"},
    "pilot_pm": {"pilot pm"},
    "pilot_exec_tl": {"pilot exec tl"},
    "integration_leader": {"integration leader"},
    "investment_type": {"investment type", "poc type"},
    "leadership": {"leadership"},
    "finance_kpi": {"finance kpi"},
    "schedule_kpi": {"schedule kpi"},
    "region": {"region"},
    "initiative_qua": {"initiativequa"},
    "initiative_qua_id": {"initiativequa: id", "initiativequa id"},
    "proposal_duration": {"proposal duration"},
    "external_created": {"created"},
}


class _Rollback(Exception):
    pass


# ---------------------------------------------------------------------------
# Value helpers
# ---------------------------------------------------------------------------
def _text(v):
    return str(v).strip() if v not in (None, "") else ""


def _person(v):
    """Strip SharePoint lookup suffixes like ';#77'."""
    return re.sub(r"\s*;#.*$", "", _text(v)).strip()


def _plain(v):
    """Flatten HTML (rich-text export) to readable plain text."""
    s = _text(v)
    if not s:
        return ""
    s = re.sub(r"<[^>]+>", " ", s)
    s = html.unescape(s)
    return re.sub(r"\s+", " ", s).strip()


def _clean_list(v):
    """A JSON array like '["A","B"]' → 'A, B'; otherwise the raw text."""
    s = _text(v)
    if s.startswith("["):
        try:
            arr = json.loads(s)
            if isinstance(arr, list):
                return ", ".join(str(x) for x in arr)
        except ValueError:
            pass
    return s


def _region(v):
    s = _text(v)
    if s.startswith("{"):
        try:
            return _text(json.loads(s).get("DisplayName"))
        except (ValueError, AttributeError):
            return ""
    return s


_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d")
_DT_FORMATS = ("%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S")


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, _date_cls):
        return v
    s = str(v).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    d = parse_date(s)
    if d:
        return d
    dt = parse_datetime(s)
    return dt.date() if dt else None


def _aware(v):
    if v in (None, ""):
        return None
    tz = timezone.get_default_timezone()
    if isinstance(v, datetime):
        return timezone.make_aware(v, tz) if timezone.is_naive(v) else v
    s = str(v).strip()
    for fmt in _DT_FORMATS:
        try:
            return timezone.make_aware(datetime.strptime(s, fmt), tz)
        except ValueError:
            continue
    d = _date(s)
    return timezone.make_aware(datetime(d.year, d.month, d.day), tz) if d else None


def _int(v):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# Reading + mapping
# ---------------------------------------------------------------------------
def _read_rows(source):
    """Return (headers, data_rows) from an .xlsx or .csv source (path/file-like)."""
    name = (getattr(source, "name", "") or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    if ext in {"csv", "txt"}:
        data = source.read()
        if isinstance(data, bytes):
            data = data.decode("utf-8-sig", errors="replace")
        rows = [list(r) for r in csv.reader(io.StringIO(data))]
    else:
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]

    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return [], []
    headers = [_text(h) for h in rows[0]]
    return headers, rows[1:]


def _resolve_mapping(headers):
    """{poc_field: column_index} by matching header synonyms."""
    norm = [h.strip().lower() for h in headers]
    mapping = {}
    for field, syns in FIELD_SYNONYMS.items():
        for i, h in enumerate(norm):
            if h in syns:
                mapping[field] = i
                break
    return mapping


def _build_defaults(get):
    """Map a field accessor ``get(field)`` to POC field values."""
    raw_status = _text(get("status"))
    return {
        "name": _text(get("name"))[:200],
        "description": _plain(get("description")),
        "status": _STATUS_MAP.get(raw_status.lower(), POC.Status.DRAFT),
        "external_status": raw_status[:50],
        "l2_wbs": _text(get("l2_wbs"))[:120],
        "initiative": _text(get("initiative"))[:200],
        "customer_segment": _clean_list(get("customer_segment"))[:120],
        "customer": _text(get("customer"))[:200],
        "leading_organization": _text(get("leading_organization"))[:200],
        "tendering_start": _date(get("tendering_start")),
        "tendering_finish": _date(get("tendering_finish")),
        "execution_start": _date(get("execution_start")),
        "execution_finish": _date(get("execution_finish")),
        "start_date": _date(get("execution_start")),
        "end_date": _date(get("execution_finish")),
        "bfo_no": _text(get("bfo_no"))[:120],
        "pilot_requestor": _person(get("pilot_requestor"))[:200],
        "opportunity_leader": _person(get("opportunity_leader"))[:200],
        "ecostruxure_lead": _person(get("ecostruxure_lead"))[:200],
        "pilot_tender_leader": _person(get("pilot_tender_leader"))[:200],
        "pilot_tender_tl": _person(get("pilot_tender_tl"))[:200],
        "pilot_pm": _person(get("pilot_pm"))[:200],
        "pilot_exec_tl": _person(get("pilot_exec_tl"))[:200],
        "integration_leader": _person(get("integration_leader"))[:200],
        "investment_type": _clean_list(get("investment_type"))[:120],
        "leadership": _text(get("leadership"))[:120],
        "finance_kpi": _text(get("finance_kpi"))[:60],
        "schedule_kpi": _text(get("schedule_kpi"))[:60],
        "region": _region(get("region"))[:120],
        "initiative_qua": _person(get("initiative_qua"))[:200],
        "initiative_qua_id": _text(get("initiative_qua_id"))[:60],
        "proposal_duration": _int(get("proposal_duration")),
        "external_created": _aware(get("external_created")),
    }


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------
def import_pocs_from_xlsx(source, user, dry_run=False):
    """Import POCs from an .xlsx or .csv source. Returns a stats dict.

    ``dry_run=True`` validates and counts without persisting (no DB rows, no
    blueprint phases/files created).
    """
    headers, data_rows = _read_rows(source)
    stats = {"created": 0, "updated": 0, "skipped": 0, "errors": []}
    if not data_rows:
        return stats
    mapping = _resolve_mapping(headers)

    def make_get(row):
        def get(field):
            i = mapping.get(field)
            return row[i] if i is not None and i < len(row) else None

        return get

    try:
        with transaction.atomic():
            for rownum, row in enumerate(data_rows, start=2):
                get = make_get(row)
                if not _text(get("name")):
                    stats["skipped"] += 1
                    continue
                try:
                    with transaction.atomic():  # per-row savepoint
                        defaults = _build_defaults(get)
                        eid = _int(get("external_id"))
                        if eid is not None:
                            obj, created = POC.objects.update_or_create(
                                external_id=eid,
                                defaults=defaults,
                                create_defaults={**defaults, "created_by": user},
                            )
                        else:
                            obj = POC.objects.create(created_by=user, **defaults)
                            created = True
                        if created:
                            stats["created"] += 1
                            if not dry_run:
                                apply_phase_templates(obj)
                        else:
                            stats["updated"] += 1
                except Exception as exc:  # noqa: BLE001 — isolate the bad row
                    stats["errors"].append(f"Row {rownum} ({_text(get('name'))}): {exc}")
            if dry_run:
                raise _Rollback
    except _Rollback:
        pass
    return stats
