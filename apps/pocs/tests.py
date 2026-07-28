"""Tests for the task state machine, the 'stuck in pending' fix, and derived
progress (Block A)."""

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from datetime import timedelta

from django.utils import timezone

from django.core.files.uploadedfile import SimpleUploadedFile

from apps.pocs.forms import TaskForm
from apps.pocs.importer import import_pocs_from_xlsx
from apps.pocs.models import (
    POC,
    FunctionalAnalysisStep,
    Phase,
    PhaseDocument,
    PhaseKind,
    POCMembership,
    Requirement,
    Task,
    Team,
    Test,
    TestValidation,
    UseCase,
)
from apps.pocs.services import delete_poc
from apps.pocs.state_machine import can_transition_task, task_allowed_statuses

User = get_user_model()


class StateMachineTests(TestCase):
    def test_valid_and_invalid_transitions(self):
        self.assertTrue(can_transition_task("pending", "in_progress"))
        self.assertTrue(can_transition_task("in_progress", "completed"))
        self.assertTrue(can_transition_task("completed", "in_progress"))  # reopen
        self.assertTrue(can_transition_task("in_progress", "pending"))  # revert
        # pending cannot jump straight to completed (must pass in_progress)
        self.assertFalse(can_transition_task("pending", "completed"))

    def test_allowed_statuses_includes_current(self):
        allowed = task_allowed_statuses("pending")
        self.assertIn("pending", allowed)  # current must be representable
        self.assertIn("in_progress", allowed)


class StatusUpdateViewTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            "a_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "a_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="T", created_by=self.admin, status="active")
        POCMembership.objects.create(
            poc=self.poc, user=self.member, role_in_poc="member"
        )
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        self.task = Task.objects.create(
            phase=self.phase, title="X", assigned_to=self.member, status="pending"
        )

    def test_member_can_leave_pending(self):
        """Regression: an assigned member could not move a task out of pending."""
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:task_status", args=[self.task.id]),
            {"status": "in_progress"},
        )
        self.assertEqual(resp.status_code, 200)
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "in_progress")

    def test_invalid_transition_is_noop(self):
        self.client.force_login(self.member)
        self.client.post(
            reverse("pocs:task_status", args=[self.task.id]), {"status": "completed"}
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "pending")  # unchanged

    def test_completing_stamps_and_phase_awaits_then_completes_on_approval(self):
        self.client.force_login(self.member)
        self.client.post(
            reverse("pocs:task_status", args=[self.task.id]), {"status": "in_progress"}
        )
        self.client.post(
            reverse("pocs:task_status", args=[self.task.id]), {"status": "completed"}
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "completed")
        self.assertIsNotNone(self.task.completed_at)
        self.assertEqual(self.task.completed_by, self.member)
        # Spec Fase 5: all work done, but the phase is NOT complete until the TL
        # approves — it stays in_progress (awaiting approval) meanwhile.
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, "in_progress")
        self.assertTrue(self.phase.awaiting_approval)
        # Admin (validator) approves → phase completes and locks.
        self.client.force_login(self.admin)
        self.client.post(reverse("pocs:phase_approve", args=[self.phase.id]))
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, "completed")
        self.assertTrue(self.phase.is_locked)


class ProgressDerivationTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            "p_admin", password="x", is_active=True, role=User.Role.ADMIN
        )

    def test_completed_poc_reads_100_even_without_items(self):
        poc = POC.objects.create(name="C", created_by=self.admin, status="completed")
        self.assertEqual(poc.progress_percent(), 100)  # no "completed but 0%"

    def test_progress_is_completed_phases_over_total(self):
        poc = POC.objects.create(name="D", created_by=self.admin, status="active")
        Phase.objects.create(poc=poc, name="P1", order=1, status="completed")
        Phase.objects.create(poc=poc, name="P2", order=2, status="in_progress")
        # 1 of 2 leaf phases completed → 50% (in_progress doesn't count).
        self.assertEqual(poc.progress_percent(), 50)

    def test_progress_counts_leaf_phases_not_parents(self):
        poc = POC.objects.create(name="E", created_by=self.admin, status="active")
        root = Phase.objects.create(poc=poc, name="Group", order=1)
        Phase.objects.create(poc=poc, parent=root, name="c1", order=1, status="completed")
        Phase.objects.create(poc=poc, parent=root, name="c2", order=2, status="pending")
        # Only the two leaves count (the parent is a grouping) → 1/2 = 50%.
        self.assertEqual(poc.progress_percent(), 50)


class RobustnessTests(TestCase):
    """Block B: validation of minimum fields, null assignee, safe render."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "b_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="T", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)

    def test_blank_title_rejected(self):
        form = TaskForm(data={"title": "   ", "status": "pending"}, poc=self.poc)
        self.assertFalse(form.is_valid())
        self.assertIn("title", form.errors)

    def test_title_is_trimmed(self):
        form = TaskForm(data={"title": "  Survey  ", "status": "pending"}, poc=self.poc)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["title"], "Survey")

    def test_phase_renders_with_unassigned_task_and_empty_date(self):
        # Tasks moved out of the phase page into the dedicated Tasks
        # workspace (spec item 4) — exercise the same null-assignee/empty-date
        # render path there instead.
        Task.objects.create(
            phase=self.phase, title="X", assigned_to=None, due_date=None
        )
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:tasks"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Unassigned")


class DeletePOCTests(TestCase):
    """Block C: cascade delete + double-confirmation guard."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "c_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "c_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="Doomed", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        Task.objects.create(phase=self.phase, title="t")
        Test.objects.create(phase=self.phase, title="te")

    def test_delete_poc_cascades(self):
        delete_poc(self.poc)
        self.assertFalse(POC.objects.filter(pk=self.poc.pk).exists())
        self.assertFalse(Phase.objects.filter(poc_id=self.poc.pk).exists())
        self.assertFalse(Task.objects.filter(phase__poc_id=self.poc.pk).exists())
        self.assertFalse(Test.objects.filter(phase__poc_id=self.poc.pk).exists())

    def test_view_requires_matching_name(self):
        self.client.force_login(self.admin)
        # Wrong name → nothing deleted.
        self.client.post(reverse("pocs:delete", args=[self.poc.pk]), {"confirm_name": "nope"})
        self.assertTrue(POC.objects.filter(pk=self.poc.pk).exists())
        # Correct name → deleted.
        self.client.post(reverse("pocs:delete", args=[self.poc.pk]), {"confirm_name": "Doomed"})
        self.assertFalse(POC.objects.filter(pk=self.poc.pk).exists())

    def test_member_cannot_delete(self):
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:delete", args=[self.poc.pk]), {"confirm_name": "Doomed"}
        )
        self.assertEqual(resp.status_code, 403)
        self.assertTrue(POC.objects.filter(pk=self.poc.pk).exists())


class FunctionalAnalysisTests(TestCase):
    """Block E: FA phases seed one document per step (title + guidance)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "e_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="FA POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(
            poc=self.poc,
            name="Functional Analysis",
            order=1,
            kind=PhaseKind.FUNCTIONAL_ANALYSIS,
        )
        self.step = FunctionalAnalysisStep.objects.create(
            title="Login flow", description="Validate the login", order=1
        )

    def test_fa_phase_seeds_sections_from_steps(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Functional Analysis")
        self.assertContains(resp, "Generate report")   # FA phases generate directly, no "Create report"
        # A PhaseDocument was created for the step (visiting the page seeds it).
        doc = PhaseDocument.objects.get(phase=self.phase, source_fa_step=self.step)
        self.assertEqual(doc.title, "Login flow")
        self.assertTrue(doc.is_fa_section)


class FunctionalAnalysisSectionKindTests(TestCase):
    """Use Cases/Requirements FA sections: admin-defined, auto-filled, locked."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "ek_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.lead = User.objects.create_user(
            "ek_lead", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="EK POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.lead, role_in_poc="lead")
        self.phase = Phase.objects.create(
            poc=self.poc, name="Functional Analysis", order=1,
            kind=PhaseKind.FUNCTIONAL_ANALYSIS,
        )
        self.req = Requirement.objects.create(
            poc=self.poc, sub_system="SCADA", req_gravity="imposes_mvp",
            req_operation="navigation", req_functional="performance",
            req_category="normal_operation", created_by=self.admin,
        )
        self.uc = UseCase.objects.create(poc=self.poc, title="Operator login", created_by=self.admin)
        self.uc.requirements.add(self.req)
        self.uc_step = FunctionalAnalysisStep.objects.create(
            title="Use Cases", order=1, section_kind="use_cases"
        )
        self.req_step = FunctionalAnalysisStep.objects.create(
            title="Requirements", order=2, section_kind="requirements"
        )
        self.other_step = FunctionalAnalysisStep.objects.create(
            title="Overview", order=3, section_kind="other"
        )

    def test_only_one_step_per_locked_kind_allowed(self):
        self.client.force_login(self.admin)
        resp = self.client.post(reverse("pocs:fa_step_create"), {
            "title": "Duplicate", "description": "", "section_kind": "use_cases",
        })
        self.assertEqual(resp.status_code, 200)  # re-rendered with error
        self.assertContains(resp, "only one is allowed")
        self.assertFalse(FunctionalAnalysisStep.objects.filter(title="Duplicate").exists())

    def test_locked_sections_seeded_and_synced_with_live_data(self):
        self.client.force_login(self.lead)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Functional Analysis structure")
        self.assertContains(resp, "New section")

        uc_doc = PhaseDocument.objects.get(phase=self.phase, source_fa_step=self.uc_step)
        req_doc = PhaseDocument.objects.get(phase=self.phase, source_fa_step=self.req_step)
        self.assertTrue(uc_doc.is_locked_section)
        self.assertIn(self.uc.title, uc_doc.content)
        self.assertIn(self.req.code, req_doc.content)

        # A new Use Case appears next time the phase is viewed (no manual edit needed).
        uc2 = UseCase.objects.create(poc=self.poc, title="Second flow", created_by=self.admin)
        self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        uc_doc.refresh_from_db()
        self.assertIn(uc2.title, uc_doc.content)

    def test_locked_section_preview_is_readonly(self):
        self.client.force_login(self.lead)
        self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))  # seed
        uc_doc = PhaseDocument.objects.get(phase=self.phase, source_fa_step=self.uc_step)

        resp = self.client.get(reverse("pocs:fa_section_preview", args=[uc_doc.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, self.uc.title)

        # Direct edit attempts don't change the locked content (field is disabled).
        resp = self.client.post(reverse("pocs:document_edit", args=[uc_doc.pk]), {
            "title": "Use Cases", "content": "HACKED",
        })
        self.assertEqual(resp.status_code, 302)
        uc_doc.refresh_from_db()
        self.assertNotIn("HACKED", uc_doc.content)

    def test_new_section_can_be_inserted_at_a_chosen_position(self):
        self.client.force_login(self.lead)
        self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))  # seed
        uc_doc = PhaseDocument.objects.get(phase=self.phase, source_fa_step=self.uc_step)

        resp = self.client.post(reverse("pocs:document_create", args=[self.phase.pk]), {
            "title": "Appendix", "content": "extra", "insert_after": uc_doc.pk,
        })
        self.assertEqual(resp.status_code, 302)
        ordered = list(self.phase.documents.order_by("order", "id").values_list("title", flat=True))
        self.assertEqual(ordered.index("Appendix"), ordered.index("Use Cases") + 1)

    def test_deleting_a_step_removes_its_seeded_section_everywhere(self):
        """Regression: a deleted step must not leave an orphaned "ghost" section
        that duplicates the section seeded for a later, differently-configured
        step of the same name."""
        self.client.force_login(self.lead)
        self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))  # seed
        uc_doc_id = PhaseDocument.objects.get(phase=self.phase, source_fa_step=self.uc_step).pk

        self.client.force_login(self.admin)
        resp = self.client.post(reverse("pocs:fa_step_delete", args=[self.uc_step.pk]))
        self.assertEqual(resp.status_code, 302)

        # The seeded section is gone too — not orphaned into a stray free document.
        self.assertFalse(PhaseDocument.objects.filter(pk=uc_doc_id).exists())
        self.assertFalse(self.phase.documents.filter(title="Use Cases").exists())

    def test_other_section_also_has_a_preview(self):
        self.client.force_login(self.lead)
        self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))  # seed
        other_doc = PhaseDocument.objects.get(phase=self.phase, source_fa_step=self.other_step)
        other_doc.content = "Some free text"
        other_doc.save()

        resp = self.client.get(reverse("pocs:fa_section_preview", args=[other_doc.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Some free text")
        self.assertNotContains(resp, "Auto-filled — not editable here")

    def test_fa_phase_report_option_generates_directly(self):
        self.client.force_login(self.lead)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertContains(resp, "Generate report")
        self.assertNotContains(resp, "Create report")


class RecordResultInheritsExpectedTests(TestCase):
    """Block F: the Record-result panel shows the Expected result as reference."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "f_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="F POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        Test.objects.create(
            phase=self.phase, title="T", expected_result="The light turns green"
        )

    def test_expected_result_is_shown_in_record_panel(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertContains(resp, "Expected result — validate against this")
        self.assertContains(resp, "The light turns green")
        # Record panel now has separate execution-status / result selects + notes.
        self.assertContains(resp, "Execution status")
        self.assertContains(resp, 'name="result"')


class TasksViewTests(TestCase):
    """Block G: role-based task views + filters."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "g_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "g_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="G POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        self.mine = Task.objects.create(phase=self.phase, title="Mine", assigned_to=self.member)
        self.other = Task.objects.create(phase=self.phase, title="Others task")
        self.overdue = Task.objects.create(
            phase=self.phase, title="Late one", assigned_to=self.member,
            due_date=timezone.localdate() - timedelta(days=3), status="in_progress",
        )

    def test_member_sees_every_task_in_their_poc(self):
        # Section 3 fix: read-only visibility for any POC member, not just
        # tasks assigned directly to them.
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:tasks"))
        self.assertContains(resp, "Mine")
        self.assertContains(resp, "Others task")

    def test_non_member_does_not_see_tasks_of_other_pocs(self):
        outsider = User.objects.create_user(
            "g_outsider", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.client.force_login(outsider)
        resp = self.client.get(reverse("pocs:tasks"))
        self.assertNotContains(resp, "Others task")

    def test_admin_is_manager_sees_all(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:tasks"))
        self.assertContains(resp, "Mine")
        self.assertContains(resp, "Others task")

    def test_overdue_filter(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:tasks") + "?f=overdue")
        self.assertContains(resp, "Late one")
        self.assertNotContains(resp, ">Mine<")  # the non-overdue assigned task hidden

    def test_status_dropdown_is_populated(self):
        # Regression: TasksView annotated can_lead/can_execute but never
        # allowed_statuses, so task_row.html's status <select> (which only
        # renders an <option> for values in task.allowed_statuses) rendered
        # with zero options here — even though the identical row worked fine
        # from the phase-detail page, which does set allowed_statuses.
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:tasks"))
        self.assertContains(resp, '<option value="in_progress"')

    def test_poc_with_zero_tasks_still_gets_an_add_task_entry(self):
        # A POC that has never had a task built no group at all here, so its
        # POC-level "+ Add Task" link (spec item 4) was unreachable — the
        # phase-level version of the same bug, one layer up.
        empty_poc = POC.objects.create(
            name="Empty POC", created_by=self.admin, status="active"
        )
        POCMembership.objects.create(
            poc=empty_poc, user=self.admin, role_in_poc=POCMembership.Role.LEAD
        )
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:tasks"))
        self.assertContains(resp, "Empty POC")
        self.assertContains(
            resp, reverse("pocs:task_create_for_poc", args=[empty_poc.pk])
        )

    def test_member_without_lead_role_gets_no_add_task_for_empty_poc(self):
        # The zero-task placeholder group is only worth showing to someone
        # who could actually use its "Add task" link.
        empty_poc = POC.objects.create(
            name="Empty POC 2", created_by=self.admin, status="active"
        )
        POCMembership.objects.create(poc=empty_poc, user=self.member, role_in_poc="member")
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:tasks"))
        self.assertNotContains(resp, "Empty POC 2")


class POCOverviewTasksWidgetTests(TestCase):
    """POC Overview tab no longer shows a Tasks widget (spec item 4 —
    minimalism: tasks only live in the dedicated Tasks workspace)."""

    def test_overview_no_longer_shows_task_widget(self):
        admin = User.objects.create_user(
            "widget_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        poc = POC.objects.create(name="Widget POC", created_by=admin, status="active")
        phase = Phase.objects.create(poc=poc, name="P", order=1)
        task = Task.objects.create(phase=phase, title="Widget task")
        self.client.force_login(admin)
        resp = self.client.get(reverse("pocs:detail", args=[poc.pk]))
        # Assert against the task row's own markup, not its title text — the
        # Audit Log tab on this same page legitimately mentions the task's
        # title in its "created ... Widget task" entry, which is unrelated
        # to the removed widget and must keep working.
        self.assertNotContains(resp, f'id="task-{task.id}"')


class POCCsvImportTests(TestCase):
    """Importing the newer 'PoC List' CSV schema (synonyms, DD/MM/YYYY, HTML)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "imp_admin", password="x", is_active=True, role=User.Role.ADMIN
        )

    def _csv(self):
        return (
            b"ID,PoC Name,PoC Status,PoC Start,PoC Finish,PoC Description,Served Segments\n"
            b'9001,My CSV PoC,Ongoing,02/12/2024,15/04/2026,<p>Hello <b>world</b></p>,"[""Auto"",""Grid""]"\n'
        )

    def test_dry_run_counts_without_writing(self):
        f = SimpleUploadedFile("PoC List.csv", self._csv(), content_type="text/csv")
        stats = import_pocs_from_xlsx(f, self.admin, dry_run=True)
        self.assertEqual(stats["created"], 1)
        self.assertEqual(stats["errors"], [])
        self.assertFalse(POC.objects.filter(external_id=9001).exists())  # rolled back

    def test_csv_import_maps_fields(self):
        f = SimpleUploadedFile("PoC List.csv", self._csv(), content_type="text/csv")
        import_pocs_from_xlsx(f, self.admin, dry_run=False)
        poc = POC.objects.get(external_id=9001)
        self.assertEqual(poc.name, "My CSV PoC")
        self.assertEqual(poc.external_status, "Ongoing")
        self.assertEqual(poc.status, "active")           # Ongoing → active
        self.assertEqual(poc.execution_start.isoformat(), "2024-12-02")  # DD/MM/YYYY
        self.assertEqual(poc.end_date.isoformat(), "2026-04-15")
        self.assertEqual(poc.description, "Hello world")  # HTML flattened
        self.assertEqual(poc.customer_segment, "Auto, Grid")  # JSON list joined


class PhaseStatusClickTests(TestCase):
    """Click-to-set phase status (manual override)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "ps_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "ps_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="PS", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="Milestone", order=1)

    def test_admin_sets_phase_status(self):
        self.client.force_login(self.admin)
        resp = self.client.post(
            reverse("pocs:phase_set_status", args=[self.phase.pk]), {"status": "completed"}
        )
        self.assertEqual(resp.status_code, 200)
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, "completed")

    def test_member_cannot_set_phase_status(self):
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:phase_set_status", args=[self.phase.pk]), {"status": "completed"}
        )
        self.assertEqual(resp.status_code, 403)
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, "pending")

    def test_completing_parent_cascades_to_subphases(self):
        child = Phase.objects.create(
            poc=self.poc, name="Sub", parent=self.phase, order=1, status="pending"
        )
        self.client.force_login(self.admin)
        self.client.post(
            reverse("pocs:phase_set_status", args=[self.phase.pk]),
            {"status": "completed"},
        )
        child.refresh_from_db()
        self.assertEqual(child.status, "completed")


class RolePermissionTests(TestCase):
    """The global role set is admin/team_member only; leadership is per-POC and
    a per-POC lead has full control inside that POC (regardless of lead_editable).
    """

    def setUp(self):
        self.admin = User.objects.create_user(
            "r_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.lead = User.objects.create_user(
            "r_lead", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.member = User.objects.create_user(
            "r_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="R", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.lead, role_in_poc="lead")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        # A phase NOT flagged lead_editable (default) — the lead must still edit it.
        self.phase = Phase.objects.create(
            poc=self.poc, name="Locked-by-old-rules", order=1, lead_editable=False
        )

    def test_no_global_poc_lead_role(self):
        self.assertEqual(
            set(User.Role.values), {"admin", "team_member"}
        )
        self.assertFalse(hasattr(User.Role, "POC_LEAD"))

    def test_lead_can_edit_any_phase_in_their_poc(self):
        from apps.core.mixins import user_can_edit_phase

        self.assertTrue(user_can_edit_phase(self.lead, self.phase))
        self.assertFalse(user_can_edit_phase(self.member, self.phase))

    def test_lead_can_set_phase_status(self):
        self.client.force_login(self.lead)
        resp = self.client.post(
            reverse("pocs:phase_set_status", args=[self.phase.pk]),
            {"status": "completed"},
        )
        self.assertEqual(resp.status_code, 200)
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, "completed")

    def test_lead_can_add_members(self):
        newbie = User.objects.create_user(
            "r_newbie", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.client.force_login(self.lead)
        resp = self.client.post(
            reverse("pocs:member_add", args=[self.poc.pk]),
            {"user_id": newbie.pk, "role": "member", "q": ""},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(
            POCMembership.objects.filter(poc=self.poc, user=newbie).exists()
        )

    def test_plain_member_cannot_add_members(self):
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:member_add", args=[self.poc.pk]),
            {"user_id": self.admin.pk, "role": "member", "q": ""},
        )
        self.assertEqual(resp.status_code, 403)


# ---------------------------------------------------------------------------
# Phase kinds: blueprint inheritance + documentation report generation
# ---------------------------------------------------------------------------
import base64
import os
import tempfile
from io import BytesIO

from django.test import override_settings

from apps.pocs.models import (
    BasePhaseDocument,
    PhaseImage,
    PhaseTemplate,
)
from apps.pocs.services import apply_phase_templates

# A 1x1 transparent PNG (local constant — not fetched remotely).
_PNG_1x1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _minimal_docx_bytes():
    """A valid empty .docx to serve as a report template in tests."""
    from docx import Document

    buf = BytesIO()
    Document().save(buf)
    return buf.getvalue()


class PhaseKindInheritanceTests(TestCase):
    """Creating a POC inherits each node's kind, base documents and FA sections."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "k_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        doc_node = PhaseTemplate.objects.create(
            name="Docs", kind=PhaseKind.DOCUMENTATION, order=1
        )
        BasePhaseDocument.objects.create(
            phase_template=doc_node, title="Intro", content="# Hello", order=1
        )
        PhaseTemplate.objects.create(
            name="FA", kind=PhaseKind.FUNCTIONAL_ANALYSIS, order=2
        )
        FunctionalAnalysisStep.objects.create(title="Step A", order=1)

    def test_inheritance_creates_documents_and_fa_sections(self):
        poc = POC.objects.create(name="K POC", created_by=self.admin, status="active")
        apply_phase_templates(poc)

        doc_phase = poc.phases.get(name="Docs")
        self.assertEqual(doc_phase.kind, PhaseKind.DOCUMENTATION)
        self.assertEqual(doc_phase.documents.count(), 1)
        self.assertEqual(doc_phase.documents.first().title, "Intro")

        fa_phase = poc.phases.get(name="FA")
        self.assertEqual(fa_phase.kind, PhaseKind.FUNCTIONAL_ANALYSIS)
        section = fa_phase.documents.get()
        self.assertTrue(section.is_fa_section)
        self.assertEqual(section.title, "Step A")


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class DocumentationReportTests(TestCase):
    """A documentation report uses the global default template and embeds images."""

    def setUp(self):
        from apps.reports.models import ReportSettings

        self.admin = User.objects.create_user(
            "d_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="D POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(
            poc=self.poc, name="Docs", order=1, kind=PhaseKind.DOCUMENTATION
        )
        # Global default template (no per-phase template attached).
        settings_row = ReportSettings.load()
        settings_row.default_template.save(
            "default.docx", SimpleUploadedFile("default.docx", _minimal_docx_bytes()), save=True
        )
        # An uploaded image, referenced from a document by its Markdown snippet.
        self.image = PhaseImage.objects.create(
            phase=self.phase,
            image=SimpleUploadedFile("pic.png", _PNG_1x1, content_type="image/png"),
        )
        PhaseDocument.objects.create(
            phase=self.phase,
            title="Section 1",
            content=f"Some text\n\n{self.image.markdown_snippet}\n\nMore text",
            order=1,
        )

    def test_report_generates_with_embedded_image(self):
        from docx import Document

        from apps.reports.generation import generate_phase_report_from_documents
        from apps.reports.models import GeneratedReport

        report = generate_phase_report_from_documents(self.phase, self.admin)
        self.assertEqual(
            report.status, GeneratedReport.Status.READY, report.error_message
        )
        self.assertTrue(report.output_file)

        report.output_file.open("rb")
        try:
            doc = Document(BytesIO(report.output_file.read()))
        finally:
            report.output_file.close()
        # The referenced image was embedded as an inline shape.
        self.assertGreaterEqual(len(doc.inline_shapes), 1)

    def test_report_endpoint_requires_membership(self):
        outsider = User.objects.create_user(
            "d_out", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.client.force_login(outsider)
        resp = self.client.post(
            reverse("reports:phase_documents", args=[self.phase.pk])
        )
        self.assertEqual(resp.status_code, 403)


class PhaseKindUIRenderTests(TestCase):
    """Smoke tests: the new blueprint/phase/settings pages render without errors."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "u_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="U POC", created_by=self.admin, status="active")
        self.doc_phase = Phase.objects.create(
            poc=self.poc, name="Docs", order=1, kind=PhaseKind.DOCUMENTATION
        )
        self.client.force_login(self.admin)

    def test_blueprint_create_form_has_kind_selector(self):
        resp = self.client.get(reverse("pocs:phase_template_create"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Type")
        self.assertContains(resp, "Functional Analysis")

    def test_documentation_phase_detail_renders(self):
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.doc_phase.pk]))
        self.assertEqual(resp.status_code, 200)
        # Attach (upload) alongside the sections list and its direct "Generate
        # report" action (Documentation phases behave like Functional Analysis:
        # sections are managed on the main page, not a separate "create" step).
        self.assertContains(resp, "Attach report")
        self.assertContains(resp, "Generate report")
        self.assertContains(resp, "New section")

    def test_report_settings_page_renders(self):
        resp = self.client.get(reverse("reports:settings"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Default report template")


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class PhaseReportUploadTests(TestCase):
    """A finished report can be attached to a phase (kept as-is, no extraction)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "up_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(
            name="UP POC", created_by=self.admin, status="active"
        )
        # Works regardless of kind (here: a Test phase).
        self.phase = Phase.objects.create(
            poc=self.poc, name="T", order=1, kind=PhaseKind.TEST
        )
        self.client.force_login(self.admin)

    def test_upload_creates_uploaded_report_and_audits(self):
        from apps.pocs.models import AuditLog
        from apps.reports.models import GeneratedReport

        resp = self.client.post(
            reverse("reports:phase_upload", args=[self.phase.pk]),
            {"report_file": SimpleUploadedFile("final.docx", _minimal_docx_bytes())},
        )
        self.assertEqual(resp.status_code, 302)
        report = GeneratedReport.objects.get(phase=self.phase)
        self.assertEqual(report.kind, GeneratedReport.Kind.UPLOADED)
        self.assertEqual(report.status, GeneratedReport.Status.READY)
        self.assertTrue(report.output_file)
        dl = self.client.get(reverse("reports:download", args=[report.pk]))
        self.assertEqual(dl.status_code, 200)
        self.assertTrue(
            AuditLog.objects.filter(poc=self.poc, action="report_uploaded").exists()
        )

    def test_disallowed_extension_rejected(self):
        from apps.reports.models import GeneratedReport

        resp = self.client.post(
            reverse("reports:phase_upload", args=[self.phase.pk]),
            {"report_file": SimpleUploadedFile("evil.exe", b"x")},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(GeneratedReport.objects.filter(phase=self.phase).exists())

    def test_image_upload_accepted_and_offers_preview(self):
        # Screenshots (SVG/PNG/JPEG/JPG) are valid evidence for an attached
        # report, not just office-document formats — and once attached, a
        # previewable file (image or PDF) should offer an in-tool "View".
        from apps.reports.models import GeneratedReport

        png_bytes = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
            "+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        resp = self.client.post(
            reverse("reports:phase_upload", args=[self.phase.pk]),
            {"report_file": SimpleUploadedFile("screenshot.png", png_bytes)},
        )
        self.assertEqual(resp.status_code, 302)
        report = GeneratedReport.objects.get(phase=self.phase)
        self.assertTrue(report.is_previewable)
        detail = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertContains(detail, "openFilePreview(")

    def test_test_phase_shows_tests_and_both_report_options(self):
        # A Test leaf phase exposes tests plus both report options (upload and
        # direct generation from the tests — no "Create report" editor step).
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Add test")
        self.assertContains(resp, "Attach report")
        self.assertContains(resp, "Generate report")
        self.assertNotContains(resp, "Create report")


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class ReportDeletionTests(TestCase):
    """Generated reports can be removed (with permission) and it's audited."""

    def setUp(self):
        from apps.reports.models import GeneratedReport

        self.admin = User.objects.create_user(
            "rd_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "rd_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(
            name="RD POC", created_by=self.admin, status="active"
        )
        POCMembership.objects.create(
            poc=self.poc, user=self.member, role_in_poc=POCMembership.Role.MEMBER
        )
        self.report = GeneratedReport.objects.create(
            kind=GeneratedReport.Kind.PHASE,
            title="A report",
            poc=self.poc,
            requested_by=self.admin,
            status=GeneratedReport.Status.READY,
        )
        self.report.output_file.save(
            "r.docx", SimpleUploadedFile("r.docx", _minimal_docx_bytes()), save=True
        )

    def test_admin_can_delete_and_it_is_audited(self):
        from apps.pocs.models import AuditLog
        from apps.reports.models import GeneratedReport

        self.client.force_login(self.admin)
        resp = self.client.post(reverse("reports:delete", args=[self.report.pk]))
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(GeneratedReport.objects.filter(pk=self.report.pk).exists())
        self.assertTrue(
            AuditLog.objects.filter(poc=self.poc, action="report_deleted").exists()
        )

    def test_plain_member_cannot_delete(self):
        from apps.reports.models import GeneratedReport

        self.client.force_login(self.member)
        resp = self.client.post(reverse("reports:delete", args=[self.report.pk]))
        self.assertEqual(resp.status_code, 403)
        self.assertTrue(GeneratedReport.objects.filter(pk=self.report.pk).exists())


class TaskOnAnyPhaseTests(TestCase):
    """Tasks can be added to any phase, including a parent with sub-phases."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "tap_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(
            name="TAP POC", created_by=self.admin, status="active"
        )
        self.parent = Phase.objects.create(poc=self.poc, name="Parent", order=1)
        self.child = Phase.objects.create(
            poc=self.poc, name="Child", parent=self.parent, order=1
        )
        self.client.force_login(self.admin)

    def test_can_create_task_on_parent_phase(self):
        self.assertFalse(self.parent.is_leaf)
        resp = self.client.post(
            reverse("pocs:task_create", args=[self.parent.pk]),
            {"title": "Parent task", "status": "pending"},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(self.parent.tasks.filter(title="Parent task").exists())

    def test_task_create_redirects_to_tasks_workspace(self):
        # Not back to the phase page (spec item 4 follow-up) — task creation
        # is driven from, and should land back on, the Tasks workspace.
        resp = self.client.post(
            reverse("pocs:task_create", args=[self.parent.pk]),
            {"title": "Redirect check", "status": "pending"},
        )
        self.assertRedirects(resp, reverse("pocs:tasks"))

    def test_task_create_for_poc_redirects_to_tasks_workspace(self):
        resp = self.client.post(
            reverse("pocs:task_create_for_poc", args=[self.poc.pk]),
            {"title": "Redirect check 2", "status": "pending", "phase": self.parent.pk},
        )
        self.assertRedirects(resp, reverse("pocs:tasks"))

    def test_task_edit_redirects_to_tasks_workspace(self):
        task = Task.objects.create(phase=self.parent, title="Editable")
        resp = self.client.post(
            reverse("pocs:task_edit", args=[task.pk]),
            {"title": "Edited", "status": "pending"},
        )
        self.assertRedirects(resp, reverse("pocs:tasks"))

    def test_task_delete_redirects_to_tasks_workspace(self):
        task = Task.objects.create(phase=self.parent, title="Deletable")
        resp = self.client.post(reverse("pocs:task_delete", args=[task.pk]))
        self.assertRedirects(resp, reverse("pocs:tasks"))

    def test_task_bulk_status_redirects_to_tasks_workspace(self):
        task = Task.objects.create(phase=self.parent, title="Bulk me")
        resp = self.client.post(
            reverse("pocs:task_bulk_status", args=[self.parent.pk]),
            {"status": "in_progress", "task_ids": [task.pk]},
        )
        self.assertRedirects(resp, reverse("pocs:tasks"))

    def test_parent_phase_detail_no_longer_shows_tasks(self):
        # Tasks are removed from the phase page (spec item 4 — minimalism);
        # "Add task" now lives only in the dedicated Tasks workspace.
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.parent.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Sub-phases")
        self.assertNotContains(resp, "Add task")

    def test_add_task_lives_in_tasks_workspace(self):
        self.parent.tasks.create(title="Workspace-visible task")
        resp = self.client.get(reverse("pocs:tasks"))
        self.assertContains(resp, "Add task")
        self.assertContains(resp, "Workspace-visible task")


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class PhaseTemplateAndZipTests(TestCase):
    """Zip uploads, downloadable phase templates, and .zip-template fallback."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "pt_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(
            name="PT POC", created_by=self.admin, status="active"
        )
        self.phase = Phase.objects.create(
            poc=self.poc, name="Docs", order=1, kind=PhaseKind.DOCUMENTATION
        )
        self.client.force_login(self.admin)

    def test_zip_upload_is_accepted(self):
        from apps.reports.models import GeneratedReport

        resp = self.client.post(
            reverse("reports:phase_upload", args=[self.phase.pk]),
            {"report_file": SimpleUploadedFile("bundle.zip", b"PK\x03\x04zip")},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(
            GeneratedReport.objects.filter(
                phase=self.phase, kind=GeneratedReport.Kind.UPLOADED
            ).exists()
        )

    def test_download_phase_template(self):
        self.phase.report_template.save(
            "tpl.docx", SimpleUploadedFile("tpl.docx", _minimal_docx_bytes()), save=True
        )
        resp = self.client.get(
            reverse("pocs:phase_template_download", args=[self.phase.pk])
        )
        self.assertEqual(resp.status_code, 200)

    def test_zip_template_falls_back_to_default_for_generation(self):
        from apps.reports.generation import generate_phase_report_from_documents
        from apps.reports.models import GeneratedReport, ReportSettings

        # Phase template is a .zip (download-only); a global .docx default exists.
        self.phase.report_template.save(
            "bundle.zip", SimpleUploadedFile("bundle.zip", b"PK\x03\x04"), save=True
        )
        ReportSettings.load().default_template.save(
            "def.docx", SimpleUploadedFile("def.docx", _minimal_docx_bytes()), save=True
        )
        PhaseDocument.objects.create(
            phase=self.phase, title="Sec", content="body", order=1
        )
        report = generate_phase_report_from_documents(self.phase, self.admin)
        self.assertEqual(
            report.status, GeneratedReport.Status.READY, report.error_message
        )

    def test_download_template_requires_membership(self):
        self.phase.report_template.save(
            "tpl.docx", SimpleUploadedFile("tpl.docx", _minimal_docx_bytes()), save=True
        )
        outsider = User.objects.create_user(
            "pt_out", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.client.force_login(outsider)
        resp = self.client.get(
            reverse("pocs:phase_template_download", args=[self.phase.pk])
        )
        self.assertEqual(resp.status_code, 403)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class TemplateDownloadFallbackTests(TestCase):
    """Download template applies to any phase/kind; falls back to the default."""

    def setUp(self):
        from apps.reports.models import ReportSettings

        self.admin = User.objects.create_user(
            "td_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(
            name="TD POC", created_by=self.admin, status="active"
        )
        ReportSettings.load().default_template.save(
            "def.docx", SimpleUploadedFile("def.docx", _minimal_docx_bytes()), save=True
        )
        self.client.force_login(self.admin)

    def test_download_falls_back_to_default_on_test_phase(self):
        # A Test phase with NO own template still offers the default for download.
        phase = Phase.objects.create(
            poc=self.poc, name="T", order=1, kind=PhaseKind.TEST
        )
        resp = self.client.get(reverse("pocs:phase_detail", args=[phase.pk]))
        self.assertContains(resp, "Download template")
        dl = self.client.get(
            reverse("pocs:phase_template_download", args=[phase.pk])
        )
        self.assertEqual(dl.status_code, 200)

    def test_download_shown_on_parent_phase_too(self):
        parent = Phase.objects.create(poc=self.poc, name="Parent", order=1)
        Phase.objects.create(poc=self.poc, name="Child", parent=parent, order=1)
        resp = self.client.get(reverse("pocs:phase_detail", args=[parent.pk]))
        self.assertContains(resp, "Download template")


class Fase3TestOutcomeTests(TestCase):
    """Fase 3a/3b: two-field outcome, mandatory notes, propose-then-validate."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "h_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.lead = User.objects.create_user(
            "h_lead", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.member = User.objects.create_user(
            "h_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="H POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.lead, role_in_poc="lead")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        self.test = Test.objects.create(
            phase=self.phase, title="T", assigned_to=self.member,
            expected_result="ok",
        )

    def _execute(self, **data):
        self.client.force_login(self.member)
        return self.client.post(
            reverse("pocs:test_execute", args=[self.test.pk]), data
        )

    def test_progress_change_applies_immediately(self):
        self._execute(execution_status="in_progress", result="")
        self.test.refresh_from_db()
        self.assertEqual(self.test.execution_status, "in_progress")
        self.assertFalse(self.test.validations.exists())

    def test_result_is_held_as_proposal_not_applied(self):
        self._execute(execution_status="test_completed", result="passed")
        self.test.refresh_from_db()
        # Test unchanged; a pending validation was created instead.
        self.assertEqual(self.test.execution_status, "not_tested")
        self.assertEqual(self.test.result, "")
        v = self.test.validations.get()
        self.assertEqual(v.status, "pending")
        self.assertEqual(v.result, "passed")

    def test_not_passed_requires_notes(self):
        self._execute(execution_status="test_completed", result="not_passed", actual_result="")
        self.assertFalse(self.test.validations.exists())  # rejected by validation

    def test_skipped_requires_notes(self):
        self._execute(execution_status="skipped", actual_result="")
        self.assertFalse(self.test.validations.exists())

    def test_lead_approves_and_outcome_applies(self):
        self._execute(execution_status="test_completed", result="passed")
        v = self.test.validations.get()
        self.client.force_login(self.lead)
        self.client.post(
            reverse("pocs:validation_decide", args=[v.pk]),
            {"decision": "approve", "comment": ""},
        )
        self.test.refresh_from_db()
        v.refresh_from_db()
        self.assertEqual(v.status, "approved")
        self.assertEqual(self.test.result, "passed")
        self.assertEqual(self.test.execution_status, "test_completed")

    def test_lead_rejects_leaves_test_unchanged(self):
        self._execute(execution_status="test_completed", result="passed")
        v = self.test.validations.get()
        self.client.force_login(self.lead)
        self.client.post(
            reverse("pocs:validation_decide", args=[v.pk]),
            {"decision": "reject", "comment": "not enough evidence"},
        )
        self.test.refresh_from_db()
        v.refresh_from_db()
        self.assertEqual(v.status, "rejected")
        self.assertEqual(self.test.result, "")

    def test_admin_terminal_change_applies_directly(self):
        self.client.force_login(self.admin)
        self.client.post(
            reverse("pocs:test_execute", args=[self.test.pk]),
            {"execution_status": "test_completed", "result": "passed"},
        )
        self.test.refresh_from_db()
        self.assertEqual(self.test.result, "passed")
        self.assertFalse(self.test.validations.exists())

    def test_poc_lead_keeps_authority_over_phase_sub_leader(self):
        # A phase sub-leader is assigned, but the POC lead still can validate.
        other = User.objects.create_user(
            "h_tl", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        POCMembership.objects.create(poc=self.poc, user=other, role_in_poc="member")
        self.phase.phase_leader = other
        self.phase.save()
        self._execute(execution_status="skipped", actual_result="n/a")
        v = self.test.validations.get()
        self.client.force_login(self.lead)  # POC lead, not the phase leader
        resp = self.client.post(
            reverse("pocs:validation_decide", args=[v.pk]),
            {"decision": "approve", "comment": ""},
        )
        self.assertEqual(resp.status_code, 302)
        v.refresh_from_db()
        self.assertEqual(v.status, "approved")

    def test_phase_sub_leader_can_also_validate(self):
        other = User.objects.create_user(
            "h_tl2", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        POCMembership.objects.create(poc=self.poc, user=other, role_in_poc="member")
        self.phase.phase_leader = other
        self.phase.save()
        self._execute(execution_status="skipped", actual_result="n/a")
        v = self.test.validations.get()
        self.client.force_login(other)
        resp = self.client.post(
            reverse("pocs:validation_decide", args=[v.pk]),
            {"decision": "approve", "comment": ""},
        )
        self.assertEqual(resp.status_code, 302)


class Fase3SpecsAndLinkTests(TestCase):
    """Fase 3c/3d: specs page, requirement/use-case create, test↔requirement."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "i_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "i_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="I POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        self.test = Test.objects.create(phase=self.phase, title="T", assigned_to=self.member)

    def test_specs_tab_visible_to_member(self):
        # The registry now lives as a tab on the POC detail page.
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:detail", args=[self.poc.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Use Cases &amp; Requirements")

    def test_specs_url_redirects_to_tab(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:specs", args=[self.poc.pk]))
        self.assertEqual(resp.status_code, 302)
        self.assertIn("tab=specs", resp.url)

    def test_member_cannot_create_requirement(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:requirement_create", args=[self.poc.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_admin_creates_requirement_with_auto_code(self):
        self.poc.l2_wbs = "6000020869"
        self.poc.save()
        self.client.force_login(self.admin)
        resp = self.client.post(
            reverse("pocs:requirement_create", args=[self.poc.pk]),
            {
                "req_gravity": "imposes_mvp",
                "req_operation": "cybersecurity",
                "req_functional": "performance",
                "req_category": "normal_operation",
            },
        )
        self.assertEqual(resp.status_code, 302)
        req = Requirement.objects.get(poc=self.poc)
        self.assertEqual(req.created_by, self.admin)
        # Code is auto-generated (no user input): POC-{wbs}-{CAT}-{OP}-{NNN}.
        self.assertEqual(req.code, "POC-6000020869-NO-CS-001")
        # Spec item 9: req_id falls back to the category abbreviation when
        # sub_system is blank (as it is here).
        self.assertEqual(req.req_id, "NO_001")

    def test_req_id_uses_sub_system_slug_and_numbers_per_poc_and_subsystem(self):
        req1 = Requirement.objects.create(
            poc=self.poc, sub_system="Fleet Integration",
            req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )
        self.assertEqual(req1.req_id, "FleetIntegration_001")
        req2 = Requirement.objects.create(
            poc=self.poc, sub_system="Fleet Integration",
            req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )
        self.assertEqual(req2.req_id, "FleetIntegration_002")
        # A different sub_system resets the counter.
        req3 = Requirement.objects.create(
            poc=self.poc, sub_system="Navigation",
            req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )
        self.assertEqual(req3.req_id, "Navigation_001")

    def test_req_id_unique_per_poc_not_global(self):
        # Spec item 9: req_id is only guaranteed unique WITHIN a POC — two
        # different POCs may legitimately share the same req_id.
        other_poc = POC.objects.create(name="Other POC", created_by=self.admin, status="active")
        req_a = Requirement.objects.create(
            poc=self.poc, sub_system="Navigation",
            req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )
        req_b = Requirement.objects.create(
            poc=other_poc, sub_system="Navigation",
            req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )
        self.assertEqual(req_a.req_id, "Navigation_001")
        self.assertEqual(req_b.req_id, "Navigation_001")

    def test_link_test_to_requirement(self):
        req = Requirement.objects.create(
            poc=self.poc, req_gravity="imposes_mvp",
            req_operation="navigation", req_functional="performance",
            req_category="normal_operation", created_by=self.admin,
        )
        self.client.force_login(self.member)
        self.client.post(
            reverse("pocs:test_link_requirements", args=[self.test.pk]),
            {"requirements": [req.pk]},
        )
        self.assertIn(req, self.test.requirements.all())

    def test_link_requirement_remove_oob_refreshes_unlinked_pool(self):
        # Spec item 8: unlinking a requirement from a test must OOB-refresh
        # the phase's "unlinked requirements" pool, not just the test row.
        req = Requirement.objects.create(
            poc=self.poc, req_gravity="imposes_mvp",
            req_operation="navigation", req_functional="performance",
            req_category="normal_operation", created_by=self.admin,
        )
        self.test.requirements.add(req)
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:test_link_requirement_remove", args=[self.test.pk, req.pk])
        )
        self.assertEqual(resp.status_code, 200)
        self.assertNotIn(req, self.test.requirements.all())
        body = resp.content.decode()
        self.assertIn('id="unlinked-requirements-panel"', body)
        self.assertIn('hx-swap-oob="true"', body)
        self.assertIn(req.req_id, body)

    def test_usecase_auto_code_and_title_unique(self):
        # First use case: code auto-generated.
        self.client.force_login(self.admin)
        self.client.post(
            reverse("pocs:usecase_create", args=[self.poc.pk]),
            {"title": "Login", "priority": "medium", "status": "draft"},
        )
        uc = UseCase.objects.get(poc=self.poc, title="Login")
        self.assertTrue(uc.code.startswith("POC-"))
        self.assertIn("-UC", uc.code)
        # Duplicate title in the same POC is rejected (re-rendered form).
        resp = self.client.post(
            reverse("pocs:usecase_create", args=[self.poc.pk]),
            {"title": "Login", "priority": "medium", "status": "draft"},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(UseCase.objects.filter(poc=self.poc, title="Login").count(), 1)


class Fase4Tests(TestCase):
    """Fase 4b (not-passed needs a requirement) + 4a Tests overview page."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "j_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "j_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="J POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        self.test = Test.objects.create(phase=self.phase, title="T", assigned_to=self.member)
        self.req = Requirement.objects.create(
            code="FO-J-SS-CS-001", poc=self.poc, req_gravity="imposes_mvp",
            req_operation="navigation", req_functional="performance",
            req_category="normal_operation", created_by=self.admin,
        )

    def test_not_passed_without_requirement_is_blocked(self):
        self.client.force_login(self.member)
        self.client.post(
            reverse("pocs:test_execute", args=[self.test.pk]),
            {"execution_status": "test_completed", "result": "not_passed",
             "actual_result": "broken"},
        )
        # No proposal created; test unchanged.
        self.assertFalse(self.test.validations.exists())
        self.test.refresh_from_db()
        self.assertEqual(self.test.result, "")

    def test_not_passed_with_requirement_creates_proposal(self):
        # Requirements are linked on the test itself (creation/edit), not
        # re-picked at execution time.
        self.test.requirements.set([self.req])
        self.client.force_login(self.member)
        self.client.post(
            reverse("pocs:test_execute", args=[self.test.pk]),
            {"execution_status": "test_completed", "result": "not_passed",
             "actual_result": "broken"},
        )
        v = self.test.validations.get()
        self.assertEqual(v.result, "not_passed")
        self.assertIn(self.req, self.test.requirements.all())

    def test_tests_overview_renders(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:tests_overview") + "?f=pending")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "T")

    def test_tests_overview_groups_by_phase_and_shows_stats(self):
        # Spec items 6 + 7: Phase sub-grouping and stat tiles.
        other_phase = Phase.objects.create(poc=self.poc, name="Other Phase", order=2)
        Test.objects.create(phase=other_phase, title="T2", result="passed")
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:tests_overview"))
        self.assertEqual(resp.status_code, 200)
        # Each phase gets its own sub-heading, linking to that phase's page.
        self.assertContains(resp, reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertContains(resp, reverse("pocs:phase_detail", args=[other_phase.pk]))
        self.assertContains(resp, other_phase.name)
        self.assertContains(resp, "Pass rate")

    def test_tests_overview_links_to_test_detail_page(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:tests_overview"))
        self.assertContains(resp, reverse("pocs:test_detail", args=[self.test.pk]))


class TestDetailPageTests(TestCase):
    """Spec item 5: a real, standalone full-page detail view for one Test."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "td_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "td_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.outsider = User.objects.create_user(
            "td_outsider", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="TD POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="TD Phase", order=1)
        self.test = Test.objects.create(
            phase=self.phase, title="Full picture test", assigned_to=self.member
        )

    def test_member_sees_full_detail(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:test_detail", args=[self.test.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Full picture test")
        self.assertContains(resp, self.phase.name)

    def test_non_member_is_forbidden(self):
        self.client.force_login(self.outsider)
        resp = self.client.get(reverse("pocs:test_detail", args=[self.test.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_phase_test_row_links_to_full_page(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertContains(resp, reverse("pocs:test_detail", args=[self.test.pk]))


class Fase4DerivedApprovalTests(TestCase):
    """User rule: requirement validated when all its tests pass/skip; a use case
    is approved when all its requirements are validated."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "k_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="K POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        self.req = Requirement.objects.create(
            poc=self.poc, req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )
        self.uc = UseCase.objects.create(poc=self.poc, title="UC", created_by=self.admin)
        self.uc.requirements.add(self.req)

    def test_derivation(self):
        # No tests → not validated / not approved.
        self.assertFalse(self.req.is_validated)
        self.assertFalse(self.uc.is_approved)
        t1 = Test.objects.create(phase=self.phase, title="t1")
        t2 = Test.objects.create(phase=self.phase, title="t2")
        t1.requirements.add(self.req)
        t2.requirements.add(self.req)
        # One passed, one still pending → not validated.
        t1.execution_status = "test_completed"; t1.result = "passed"; t1.save()
        self.assertFalse(self.req.is_validated)
        # Second skipped → all settled/ok → validated → use case approved.
        t2.execution_status = "skipped"; t2.actual_result = "n/a"; t2.save()
        self.assertTrue(self.req.is_validated)
        self.assertTrue(self.uc.is_approved)
        # A not-passed test breaks validation.
        t2.execution_status = "test_completed"; t2.result = "not_passed"; t2.save()
        self.assertFalse(self.req.is_validated)
        self.assertFalse(self.uc.is_approved)

    def test_requirement_links_use_case_from_requirement_form(self):
        self.client.force_login(self.admin)
        resp = self.client.post(
            reverse("pocs:requirement_create", args=[self.poc.pk]),
            {
                "req_gravity": "imposes_mvp", "req_operation": "navigation",
                "req_functional": "performance", "req_category": "maintenance_mode",
                "use_cases": [self.uc.pk],
            },
        )
        self.assertEqual(resp.status_code, 302)
        new_req = Requirement.objects.exclude(pk=self.req.pk).get(poc=self.poc)
        self.assertIn(self.uc, new_req.use_cases.all())


class Fase5PhaseApprovalTests(TestCase):
    """Fase 5: TL approval gates phase completion + locks it for editing."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "m_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.lead = User.objects.create_user(
            "m_lead", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.tl = User.objects.create_user(
            "m_tl", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.member = User.objects.create_user(
            "m_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="M POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.lead, role_in_poc="lead")
        POCMembership.objects.create(poc=self.poc, user=self.tl, role_in_poc="member")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1, phase_leader=self.tl)
        self.task = Task.objects.create(phase=self.phase, title="t", assigned_to=self.member)

    def test_cannot_approve_with_unfinished_work(self):
        self.client.force_login(self.tl)
        self.client.post(reverse("pocs:phase_approve", args=[self.phase.id]))
        self.phase.refresh_from_db()
        self.assertFalse(self.phase.is_approved)

    def test_phase_leader_approves_and_locks(self):
        self.task.status = "completed"; self.task.save()
        self.phase.recalculate_status()
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, "in_progress")  # awaiting approval
        self.client.force_login(self.tl)
        self.client.post(reverse("pocs:phase_approve", args=[self.phase.id]))
        self.phase.refresh_from_db()
        self.assertTrue(self.phase.is_approved)
        self.assertEqual(self.phase.status, "completed")

    def test_non_validator_cannot_approve(self):
        self.task.status = "completed"; self.task.save()
        self.phase.recalculate_status()
        self.client.force_login(self.member)
        resp = self.client.post(reverse("pocs:phase_approve", args=[self.phase.id]))
        self.assertEqual(resp.status_code, 403)

    def test_locked_phase_blocks_task_edits(self):
        self.task.status = "completed"; self.task.save()
        self.phase.recalculate_status()
        self.phase.approved_at = timezone.now(); self.phase.approved_by = self.tl
        self.phase.save()
        # Assigned member can no longer change the task while locked.
        self.client.force_login(self.member)
        self.client.post(
            reverse("pocs:task_status", args=[self.task.id]), {"status": "in_progress"}
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "completed")  # unchanged (locked)

    def test_poc_lead_can_unlock(self):
        self.phase.approved_at = timezone.now(); self.phase.approved_by = self.tl
        self.phase.save()
        self.client.force_login(self.lead)  # POC lead keeps authority
        self.client.post(reverse("pocs:phase_unlock", args=[self.phase.id]))
        self.phase.refresh_from_db()
        self.assertFalse(self.phase.is_approved)


class Fase6ClosureTests(TestCase):
    """Fase 6: official closure locks the POC + generates a sealed final report."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "n_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.lead = User.objects.create_user(
            "n_lead", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.member = User.objects.create_user(
            "n_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="N POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.lead, role_in_poc="lead")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        self.test = Test.objects.create(phase=self.phase, title="T", assigned_to=self.member)

    def test_member_cannot_close(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:close", args=[self.poc.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_lead_closes_and_final_report_generated(self):
        self.client.force_login(self.lead)
        resp = self.client.post(
            reverse("pocs:close", args=[self.poc.pk]),
            {"closure_date": "2026-07-02", "closure_conclusion": "All good.", "confirm": "on"},
        )
        self.assertEqual(resp.status_code, 302)
        self.poc.refresh_from_db()
        self.assertTrue(self.poc.is_closed)
        self.assertEqual(self.poc.closed_by, self.lead)
        # A sealed final report was generated for the POC.
        fr = self.poc.generated_reports.filter(kind="final").first()
        self.assertIsNotNone(fr)
        self.assertEqual(fr.status, "ready")

    def test_closed_poc_blocks_editing(self):
        self.poc.closed_at = timezone.now(); self.poc.closed_by = self.lead
        self.poc.closure_date = "2026-07-02"; self.poc.save()
        # Member can't execute the test anymore.
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:test_execute", args=[self.test.pk]),
            {"execution_status": "in_progress"},
        )
        self.assertEqual(resp.status_code, 403)
        # Lead can't edit the phase.
        self.client.force_login(self.lead)
        resp = self.client.get(reverse("pocs:phase_edit", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_only_admin_can_reopen(self):
        self.poc.closed_at = timezone.now(); self.poc.closed_by = self.lead
        self.poc.closure_date = "2026-07-02"; self.poc.save()
        self.client.force_login(self.lead)
        resp = self.client.post(reverse("pocs:reopen", args=[self.poc.pk]))
        self.assertEqual(resp.status_code, 403)
        self.client.force_login(self.admin)
        self.client.post(reverse("pocs:reopen", args=[self.poc.pk]))
        self.poc.refresh_from_db()
        self.assertFalse(self.poc.is_closed)

    def test_late_report_is_signed(self):
        from apps.reports.models import GeneratedReport
        self.poc.closed_at = timezone.now(); self.poc.closed_by = self.lead
        self.poc.closure_date = "2026-07-02"; self.poc.save()
        self.client.force_login(self.member)
        upload = SimpleUploadedFile("late.txt", b"late doc", content_type="text/plain")
        self.client.post(
            reverse("reports:phase_upload", args=[self.phase.pk]), {"report_file": upload}
        )
        r = GeneratedReport.objects.filter(poc=self.poc, kind="uploaded").first()
        self.assertIsNotNone(r)
        self.assertTrue(r.after_closure)

    def test_final_report_markdown_flags_pending_phases(self):
        from apps.reports.generation import build_final_report_markdown
        md = build_final_report_markdown(self.poc, "Conclusion text")
        self.assertIn("Final Report", md)
        self.assertIn("Conclusion text", md)
        # The phase has an unfinished test → flagged pending.
        self.assertIn("pending", md.lower())


class TestReportMarkdownTests(TestCase):
    """Spec items 2 + 3: ID+Title chapter headings, Parameters integrity."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "trm_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="TRM POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(
            poc=self.poc, name="Unit Testing", order=1, kind=PhaseKind.TEST
        )

    def test_heading_includes_id_and_title(self):
        from apps.reports.generation import build_phase_body_markdown

        Test.objects.create(phase=self.phase, title="Login works", test_code="UT-001")
        md = build_phase_body_markdown(self.phase)
        self.assertIn("UT-001", md)
        self.assertIn("UT-001 — Login works", md)

    def test_expected_result_and_target_date_always_shown(self):
        from apps.reports.generation import build_phase_body_markdown

        Test.objects.create(phase=self.phase, title="No expected result")
        md = build_phase_body_markdown(self.phase)
        self.assertIn("**Expected result:**", md)
        self.assertIn("_Not defined._", md)
        self.assertIn("**Target date:**", md)


class Fase7TaskOrderingTests(TestCase):
    """Fase 7: due-date ordering (completed last)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "o_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="O POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)

    def test_order_tasks_due_date_completed_last(self):
        from apps.pocs.views import order_tasks
        today = timezone.localdate()
        late = Task.objects.create(phase=self.phase, title="late", due_date=today + timedelta(days=5))
        soon = Task.objects.create(phase=self.phase, title="soon", due_date=today + timedelta(days=1))
        nodate = Task.objects.create(phase=self.phase, title="nodate")
        done = Task.objects.create(phase=self.phase, title="done", due_date=today, status="completed")
        ordered = list(order_tasks(self.phase.tasks.all()))
        # soon (earliest due) → late → nodate (null last) → done (completed last)
        self.assertEqual(ordered, [soon, late, nodate, done])


class TaskGanttTests(TestCase):
    """Section 5: the tree+Gantt hybrid — a summary (parent) row aggregates
    its sub-tasks' dates/progress instead of showing its own, and every row
    (dated or not) still appears in the tree."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "gt_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="GT POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)

    def test_leaf_task_uses_its_own_dates(self):
        from apps.pocs.gantt import build_task_gantt
        today = timezone.localdate()
        t = Task.objects.create(
            phase=self.phase, title="Solo", start_date=today, due_date=today + timedelta(days=3)
        )
        gantt = build_task_gantt([t], today)
        self.assertTrue(gantt["has_dates"])
        row = gantt["top_level"][0]
        self.assertFalse(row.is_summary)
        self.assertEqual(row.gantt_start, today)
        self.assertEqual(row.gantt_end, today + timedelta(days=3))
        self.assertEqual(row.gantt_progress, 0)

    def test_summary_task_aggregates_subtasks(self):
        from apps.pocs.gantt import build_task_gantt
        today = timezone.localdate()
        parent = Task.objects.create(
            phase=self.phase, title="Parent",
            start_date=today + timedelta(days=10), due_date=today + timedelta(days=10),
        )
        Task.objects.create(
            phase=self.phase, title="Sub1", parent=parent,
            start_date=today, due_date=today + timedelta(days=2), status="completed",
        )
        Task.objects.create(
            phase=self.phase, title="Sub2", parent=parent,
            start_date=today + timedelta(days=1), due_date=today + timedelta(days=5),
        )
        tasks = list(Task.objects.filter(phase=self.phase))
        gantt = build_task_gantt(tasks, today)
        self.assertEqual(len(gantt["top_level"]), 1)
        parent_row = gantt["top_level"][0]
        self.assertTrue(parent_row.is_summary)
        # Spans the earliest sub-task start to the latest sub-task end —
        # NOT the parent's own (irrelevant) start/due date.
        self.assertEqual(parent_row.gantt_start, today)
        self.assertEqual(parent_row.gantt_end, today + timedelta(days=5))
        self.assertEqual(parent_row.gantt_progress, 50)  # 1 of 2 sub-tasks done
        self.assertEqual(len(parent_row.child_list), 2)

    def test_undated_tasks_still_appear_without_a_bar(self):
        from apps.pocs.gantt import build_task_gantt
        t = Task.objects.create(phase=self.phase, title="No dates")
        gantt = build_task_gantt([t], timezone.localdate())
        self.assertFalse(gantt["has_dates"])
        self.assertEqual(gantt["top_level"], [t])
        self.assertIsNone(t.pct_start)

    def test_gantt_summary_flattens_counts_and_range(self):
        from apps.pocs.gantt import build_task_gantt, gantt_summary
        today = timezone.localdate()
        Task.objects.create(
            phase=self.phase, title="A", start_date=today, due_date=today + timedelta(days=1),
            status="completed",
        )
        Task.objects.create(
            phase=self.phase, title="B", start_date=today + timedelta(days=2),
            due_date=today + timedelta(days=4),
        )
        tasks = list(Task.objects.filter(phase=self.phase))
        gantt = build_task_gantt(tasks, today)
        summary = gantt_summary(gantt["top_level"])
        self.assertEqual(summary["total"], 2)
        self.assertEqual(summary["completed"], 1)
        self.assertEqual(summary["start"], today)
        self.assertEqual(summary["end"], today + timedelta(days=4))


class Fase8BlueprintTests(TestCase):
    """Fase 8: blueprint version history (undo/restore) + apply to a subset."""

    def setUp(self):
        from apps.pocs.models import PhaseTemplate
        self.admin = User.objects.create_user(
            "q_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.n1 = PhaseTemplate.objects.create(name="Root", order=1)

    def test_snapshot_and_restore_recreates_deleted_node(self):
        from apps.pocs.models import PhaseTemplate
        from apps.pocs.services import snapshot_blueprint, restore_blueprint
        nid = self.n1.id
        v = snapshot_blueprint(self.admin, "before delete")
        self.n1.delete()  # Django sets self.n1.pk = None after delete
        self.assertEqual(PhaseTemplate.objects.count(), 0)
        restore_blueprint(v)
        # Node restored WITH its original id (so POC links survive).
        self.assertTrue(PhaseTemplate.objects.filter(id=nid, name="Root").exists())

    def test_undo_view_reverts_last_change(self):
        from apps.pocs.models import PhaseTemplate, BlueprintVersion
        self.client.force_login(self.admin)
        # Create a node via the view (snapshots pre-state).
        self.client.post(reverse("pocs:phase_template_create"), {"name": "New", "kind": "test"})
        self.assertTrue(PhaseTemplate.objects.filter(name="New").exists())
        self.assertTrue(BlueprintVersion.objects.exists())
        # Undo → the "New" node disappears (restored to pre-create snapshot).
        self.client.post(reverse("pocs:blueprint_undo"))
        self.assertFalse(PhaseTemplate.objects.filter(name="New").exists())

    def test_apply_to_selected_pocs_only(self):
        from apps.pocs.models import Phase
        a = POC.objects.create(name="A", created_by=self.admin, status="active")
        b = POC.objects.create(name="B", created_by=self.admin, status="active")
        self.client.force_login(self.admin)
        self.client.post(
            reverse("pocs:phase_template_apply_selected"), {"poc_ids": [a.id]}
        )
        self.assertTrue(Phase.objects.filter(poc=a, source_template=self.n1).exists())
        self.assertFalse(Phase.objects.filter(poc=b, source_template=self.n1).exists())


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class Fase11MissingBlueprintTemplateTests(TestCase):
    """A blueprint node whose report_template file is missing on disk (e.g.
    media not carried over on deploy) must not crash 'Apply to all'."""

    def setUp(self):
        from apps.pocs.models import PhaseTemplate

        self.admin = User.objects.create_user(
            "mbt_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.node = PhaseTemplate.objects.create(name="FA", order=1, kind=PhaseKind.FUNCTIONAL_ANALYSIS)
        self.node.report_template.save(
            "tpl.docx", SimpleUploadedFile("tpl.docx", _minimal_docx_bytes()), save=True
        )
        # Simulate the file being absent from disk while the DB still points at it.
        import os
        os.remove(self.node.report_template.path)
        self.poc = POC.objects.create(name="MBT POC", created_by=self.admin, status="active")

    def test_apply_all_survives_missing_file(self):
        from apps.pocs.services import sync_blueprint_to_pocs

        stats = sync_blueprint_to_pocs()
        self.assertEqual(stats["created"], 1)
        self.assertIn("FA", stats["missing_templates"])
        phase = self.poc.phases.get(source_template=self.node)
        self.assertFalse(phase.report_template)

    def test_apply_all_view_warns_instead_of_500(self):
        self.client.force_login(self.admin)
        resp = self.client.post(reverse("pocs:phase_template_apply_all"), follow=True)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "missing on disk")


class Fase10Tests(TestCase):
    """Fase 10: status board (10a) + requirement Excel import (10c)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "r_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "r_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="R POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="System Testing", order=1)

    def test_overview_does_not_show_removed_status_board(self):
        t = Test.objects.create(phase=self.phase, title="t1")
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:detail", args=[self.poc.pk]))
        self.assertNotContains(resp, "Test status by system")

    def _xlsx(self, rows):
        import io, openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["sub_system", "req_gravity", "req_operation", "req_functional", "req_category", "remarks"])
        for r in rows:
            ws.append(r)
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        buf.name = "reqs.xlsx"
        return buf

    def test_import_preview_and_confirm(self):
        self.client.force_login(self.admin)
        f = self._xlsx([
            ["SCADA", "Imposes (MVP)", "Cybersecurity", "Performance", "Normal Operation", "ok"],
            ["HMI", "Custom Gravity", "navigation", "performance", "normal_operation", "new custom value"],
            ["PLC", "", "navigation", "performance", "normal_operation", "missing gravity"],
        ])
        # Stage 1: upload → preview (2 valid — a never-seen-before gravity is
        # accepted as a new value —, 1 error for the missing required field).
        resp = self.client.post(reverse("pocs:requirement_import", args=[self.poc.pk]), {"file": f})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "2 valid")
        self.assertContains(resp, "Gravity is required")
        # Stage 2: confirm → only the valid rows are created.
        resp = self.client.post(reverse("pocs:requirement_import", args=[self.poc.pk]), {"confirm": "1"})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(self.poc.requirements.count(), 2)
        req = self.poc.requirements.get(sub_system="SCADA")
        self.assertEqual(req.req_operation, "cybersecurity")  # label resolved
        self.assertTrue(req.code.startswith("POC-"))  # auto code
        custom_req = self.poc.requirements.get(sub_system="HMI")
        self.assertEqual(custom_req.req_gravity, "Custom Gravity")  # new value accepted as-is
        # ...and registered so it becomes a suggestion for this POC going forward.
        from apps.pocs.models import RequirementFieldOption

        self.assertTrue(
            RequirementFieldOption.objects.filter(
                poc=self.poc, field="req_gravity", value="Custom Gravity"
            ).exists()
        )

    def test_import_links_use_cases_and_defaults_to_none(self):
        uc = UseCase.objects.create(poc=self.poc, title="Login", created_by=self.admin)
        import io, openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([
            "sub_system", "req_gravity", "req_operation", "req_functional",
            "req_category", "remarks", "use_cases",
        ])
        ws.append(["SCADA", "imposes_mvp", "navigation", "performance", "normal_operation", "linked", uc.code])
        ws.append(["HMI", "imposes_mvp", "navigation", "performance", "normal_operation", "no link", ""])
        ws.append(["PLC", "imposes_mvp", "navigation", "performance", "normal_operation", "bad code", "NOPE-001"])
        buf = io.BytesIO(); wb.save(buf); buf.seek(0); buf.name = "reqs.xlsx"

        self.client.force_login(self.admin)
        resp = self.client.post(reverse("pocs:requirement_import", args=[self.poc.pk]), {"file": buf})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "2 valid")
        self.assertContains(resp, "Unknown use case code")
        resp = self.client.post(reverse("pocs:requirement_import", args=[self.poc.pk]), {"confirm": "1"})
        self.assertEqual(resp.status_code, 302)
        linked_req = self.poc.requirements.get(sub_system="SCADA")
        self.assertIn(uc, linked_req.use_cases.all())
        # Blank / omitted use_cases → no use case linked (previous, default behaviour).
        unlinked_req = self.poc.requirements.get(sub_system="HMI")
        self.assertFalse(unlinked_req.use_cases.exists())

    def test_import_rejects_duplicate_validation_criteria(self):
        # An existing requirement already carries this validation criteria.
        Requirement.objects.create(
            poc=self.poc, req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            validation_criteria="Response time under 200ms", created_by=self.admin,
        )
        import io, openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([
            "sub_system", "req_gravity", "req_operation", "req_functional",
            "req_category", "validation_criteria",
        ])
        # Row 1: duplicates the existing requirement above.
        ws.append(["SCADA", "imposes_mvp", "navigation", "performance", "normal_operation", "Response time under 200ms"])
        # Row 2: unique — valid.
        ws.append(["HMI", "imposes_mvp", "navigation", "performance", "normal_operation", "Screen loads under 1s"])
        # Row 3: duplicates row 2 within the same file.
        ws.append(["PLC", "imposes_mvp", "navigation", "performance", "normal_operation", "Screen loads under 1s"])
        buf = io.BytesIO(); wb.save(buf); buf.seek(0); buf.name = "reqs.xlsx"

        self.client.force_login(self.admin)
        resp = self.client.post(reverse("pocs:requirement_import", args=[self.poc.pk]), {"file": buf})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "1 valid")
        self.assertContains(resp, "Duplicate requirement")
        resp = self.client.post(reverse("pocs:requirement_import", args=[self.poc.pk]), {"confirm": "1"})
        self.assertEqual(resp.status_code, 302)
        # Only the unique row was imported — the pre-existing requirement is untouched
        # and no duplicate of it (nor the in-file repeat) was created.
        self.assertEqual(
            self.poc.requirements.filter(validation_criteria="Response time under 200ms").count(), 1
        )
        self.assertEqual(
            self.poc.requirements.filter(validation_criteria="Screen loads under 1s").count(), 1
        )

    def test_member_cannot_import(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:requirement_import", args=[self.poc.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_guidelines_shown_on_specs_tab(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:detail", args=[self.poc.pk]))
        self.assertContains(resp, "Guidelines")
        self.assertContains(resp, "Imposes (MVP)")


class UseCaseImportAndDashboardTests(TestCase):
    """Use case Excel import + dashboard shows all POCs (no pagination)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "s_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "s_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="S POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.req = Requirement.objects.create(
            poc=self.poc, sub_system="SCADA", req_gravity="imposes_mvp",
            req_operation="navigation", req_functional="performance",
            req_category="normal_operation", created_by=self.admin,
        )

    def _xlsx(self, rows):
        import io, openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["title", "priority", "status", "requirements", "remarks"])
        for r in rows:
            ws.append(r)
        buf = io.BytesIO(); wb.save(buf); buf.seek(0); buf.name = "ucs.xlsx"
        return buf

    def test_usecase_import_preview_and_confirm(self):
        self.client.force_login(self.admin)
        f = self._xlsx([
            ["Login", "High", "Active", self.req.code, "ok"],
            ["", "medium", "draft", "", "missing title"],
            ["BadReq", "medium", "draft", "NOPE-001", "unknown req"],
        ])
        resp = self.client.post(reverse("pocs:usecase_import", args=[self.poc.pk]), {"file": f})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "1 valid")
        self.assertContains(resp, "Title is required")
        self.assertContains(resp, "Unknown requirement code")
        resp = self.client.post(reverse("pocs:usecase_import", args=[self.poc.pk]), {"confirm": "1"})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(self.poc.use_cases.count(), 1)
        uc = self.poc.use_cases.get()
        self.assertEqual(uc.title, "Login")
        self.assertEqual(uc.priority, "high")   # label resolved
        self.assertIn(self.req, uc.requirements.all())
        self.assertTrue(uc.code.startswith("POC-"))

    def test_member_cannot_import_usecases(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:usecase_import", args=[self.poc.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_dashboard_shows_all_pocs_no_pagination(self):
        for i in range(15):
            POC.objects.create(name=f"P{i}", created_by=self.admin, status="active")
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("core:dashboard"))
        self.assertEqual(resp.status_code, 200)
        # All 16 POCs (S POC + 15) present; no pagination context.
        self.assertNotIn("page_obj", resp.context)
        self.assertContains(resp, "P14")
        self.assertContains(resp, "P0")


from django.test import override_settings
import tempfile

from apps.pocs.models import EvidenceFile, POCImage


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class Fase11EvidenceAndStatusTests(TestCase):
    """Multiple evidence files, and the status badge during pending validation."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "f11_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "f11_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="F11 POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="Unit Testing", order=1)
        self.test = Test.objects.create(phase=self.phase, title="T", assigned_to=self.member)

    def test_multiple_evidence_files_held_pending_then_reparented(self):
        self.client.force_login(self.member)
        img = SimpleUploadedFile("pic.png", _PNG_1x1, content_type="image/png")
        doc = SimpleUploadedFile("sheet.xlsx", b"fake-xlsx-bytes")
        resp = self.client.post(
            reverse("pocs:test_execute", args=[self.test.pk]),
            {
                "execution_status": "test_completed",
                "result": "passed",
                "actual_result": "looks fine",
                "evidence_files": [img, doc],
            },
        )
        self.assertEqual(resp.status_code, 200)
        validation = self.test.validations.get()
        self.assertEqual(validation.evidence_files.count(), 2)
        # Test itself untouched while pending.
        self.assertEqual(self.test.evidence_files.count(), 0)
        # Badge shows the PROPOSED outcome, not "not tested".
        self.test.refresh_from_db()
        self.assertEqual(self.test.display_execution_status, "test_completed")
        self.assertEqual(self.test.display_result, "passed")
        self.assertEqual(self.test.execution_status, "not_tested")

        # Admin approves → evidence re-parents onto the Test, none left on the validation.
        self.client.force_login(self.admin)
        self.client.post(reverse("pocs:validation_decide", args=[validation.pk]), {"decision": "approve"})
        self.test.refresh_from_db()
        self.assertEqual(self.test.evidence_files.count(), 2)
        validation.refresh_from_db()
        self.assertEqual(validation.evidence_files.count(), 0)

    def test_rejected_validation_deletes_its_evidence(self):
        self.client.force_login(self.member)
        img = SimpleUploadedFile("pic.png", _PNG_1x1, content_type="image/png")
        self.client.post(
            reverse("pocs:test_execute", args=[self.test.pk]),
            {
                "execution_status": "test_completed",
                "result": "passed",
                "actual_result": "looks fine",
                "evidence_files": [img],
            },
        )
        validation = self.test.validations.get()
        self.assertEqual(validation.evidence_files.count(), 1)
        self.client.force_login(self.admin)
        self.client.post(
            reverse("pocs:validation_decide", args=[validation.pk]),
            {"decision": "reject", "comment": "no good"},
        )
        self.assertEqual(EvidenceFile.objects.count(), 0)

    def test_bad_evidence_extension_rejected(self):
        self.client.force_login(self.member)
        bad = SimpleUploadedFile("virus.exe", b"x")
        resp = self.client.post(
            reverse("pocs:test_execute", args=[self.test.pk]),
            {"execution_status": "in_progress", "evidence_files": [bad]},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(EvidenceFile.objects.count(), 0)

    def test_admin_direct_apply_keeps_all_evidence_files(self):
        # Admin submissions apply immediately (no pending validation) — every
        # uploaded file must still be attached, not just the first.
        self.client.force_login(self.admin)
        f1 = SimpleUploadedFile("a.png", _PNG_1x1, content_type="image/png")
        f2 = SimpleUploadedFile("b.png", _PNG_1x1, content_type="image/png")
        resp = self.client.post(
            reverse("pocs:test_execute", args=[self.test.pk]),
            {
                "execution_status": "test_completed",
                "result": "passed",
                "actual_result": "looks fine",
                "evidence_files": [f1, f2],
            },
        )
        self.assertEqual(resp.status_code, 200)
        self.test.refresh_from_db()
        self.assertEqual(self.test.evidence_files.count(), 2)


class Fase11TestTimelineTests(TestCase):
    """Testing roadmap: rows per Test phase, tests plotted chronologically."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "f11t_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="F11T POC", created_by=self.admin, status="active")
        self.ut = Phase.objects.create(poc=self.poc, name="Unitary Test", order=1, kind=PhaseKind.TEST)
        self.st = Phase.objects.create(poc=self.poc, name="System Test", order=2, kind=PhaseKind.TEST)

    def test_no_tests_returns_none(self):
        self.assertIsNone(self.poc.test_timeline())

    def test_row_status_and_marker_positions(self):
        Test.objects.create(
            phase=self.ut, title="ok", execution_status="test_completed", result="passed",
            executed_at=timezone.now(),
        )
        Test.objects.create(phase=self.st, title="pending")
        timeline = self.poc.test_timeline()
        self.assertIsNotNone(timeline)
        rows_by_phase = {r["phase"].name: r for r in timeline["rows"]}
        self.assertEqual(rows_by_phase["Unitary Test"]["status"], "as_expected")
        self.assertEqual(rows_by_phase["System Test"]["status"], "not_started")
        # A pending test still gets a plottable position (no crash / KeyError).
        pending_marker = rows_by_phase["System Test"]["markers"][0]
        self.assertIsInstance(pending_marker["pos"], int)
        self.assertFalse(timeline["all_complete"])

    def test_all_complete_when_every_row_passed(self):
        Test.objects.create(
            phase=self.ut, title="ok", execution_status="test_completed", result="passed",
            executed_at=timezone.now(),
        )
        Test.objects.create(
            phase=self.st, title="ok2", execution_status="test_completed", result="passed",
            executed_at=timezone.now(),
        )
        timeline = self.poc.test_timeline()
        self.assertTrue(timeline["all_complete"])

    def test_overview_renders_timeline(self):
        Test.objects.create(phase=self.ut, title="ok")
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:detail", args=[self.poc.pk]))
        self.assertContains(resp, "Testing roadmap")


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class Fase11FinalReportStructureTests(TestCase):
    """Final report = Functional Analysis + Testing and Validation + Conclusions only."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "f11r_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="F11R POC", created_by=self.admin, status="active")
        self.fa_phase = Phase.objects.create(
            poc=self.poc, name="Functional Analysis", order=1, kind=PhaseKind.FUNCTIONAL_ANALYSIS
        )
        PhaseDocument.objects.create(phase=self.fa_phase, title="Overview", content="FA content here.")
        self.test_phase = Phase.objects.create(
            poc=self.poc, name="Unit Testing", order=2, kind=PhaseKind.TEST
        )
        Test.objects.create(
            phase=self.test_phase, title="Login works",
            execution_status="test_completed", result="passed",
        )
        self.doc_phase = Phase.objects.create(
            poc=self.poc, name="Docs", order=3, kind=PhaseKind.DOCUMENTATION
        )
        PhaseDocument.objects.create(phase=self.doc_phase, title="ReadMe", content="Should NOT appear.")

    def test_final_report_markdown_structure(self):
        from apps.reports.generation import build_final_report_markdown

        md = build_final_report_markdown(self.poc, conclusion="Great pilot.")
        # Chapter order.
        fa_idx = md.index("# Functional Analysis")
        tv_idx = md.index("# Testing and Validation")
        concl_idx = md.index("# Conclusions")
        self.assertTrue(fa_idx < tv_idx < concl_idx)
        # Content landed under the right chapter.
        self.assertIn("Overview", md[fa_idx:tv_idx])
        self.assertIn("FA content here.", md[fa_idx:tv_idx])
        self.assertIn("Login works", md[tv_idx:concl_idx])
        self.assertIn("Great pilot.", md[concl_idx:])
        # Documentation-kind phase content is excluded entirely.
        self.assertNotIn("Should NOT appear.", md)
        self.assertNotIn("ReadMe", md)

    def test_final_report_generates_as_docx(self):
        from apps.reports.generation import generate_final_report
        from apps.reports.models import GeneratedReport

        report = generate_final_report(self.poc, self.admin, conclusion="Done.")
        self.assertEqual(report.status, GeneratedReport.Status.READY, report.error_message)
        self.assertTrue(report.output_file)


class PocGraphStatusTests(TestCase):
    """POC node colour now derives from its Use Cases (🎉 Improvements #2):
    green only when active AND every use case is green; orange otherwise."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "g_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="G POC", created_by=self.admin, status="draft")

    def test_draft_poc_is_orange_regardless_of_use_cases(self):
        self.assertEqual(self.poc.graph_status(), "orange")

    def test_active_poc_with_no_use_cases_is_orange(self):
        self.poc.status = "active"
        self.poc.save()
        self.assertEqual(self.poc.graph_status(), "orange")

    def test_active_poc_orange_until_all_use_cases_green(self):
        self.poc.status = "active"
        self.poc.save()
        uc1 = UseCase.objects.create(poc=self.poc, title="UC1", created_by=self.admin)
        UseCase.objects.create(poc=self.poc, title="UC2", created_by=self.admin)
        # uc1 has no requirements yet -> gray, so the POC is not all-green.
        self.assertEqual(uc1.graph_status(), "gray")
        self.assertEqual(self.poc.graph_status(), "orange")

    def test_active_poc_green_only_when_every_use_case_green(self):
        self.poc.status = "active"
        self.poc.save()
        phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        req = Requirement.objects.create(
            poc=self.poc, req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )
        uc = UseCase.objects.create(poc=self.poc, title="UC", created_by=self.admin)
        uc.requirements.add(req)
        test = Test.objects.create(phase=phase, title="t")
        test.requirements.add(req)
        # Not passed yet -> use case not green -> POC stays orange.
        self.assertEqual(self.poc.graph_status(), "orange")
        test.execution_status = "test_completed"
        test.result = "passed"
        test.save()
        self.assertEqual(uc.graph_status(), "green")
        self.assertEqual(self.poc.graph_status(), "green")
        # A second, unfinished use case pulls the POC back to orange.
        UseCase.objects.create(poc=self.poc, title="UC2", created_by=self.admin)
        self.assertEqual(self.poc.graph_status(), "orange")


class CommentFeedbackTests(TestCase):
    """Improvement brief item 3: any POC member opens a thread and may reply;
    only a POC lead/admin may Close it (no separate Applied/Rejected verdict).
    An open thread shows as a halo on its target until closed."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "c_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "c_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.lead = User.objects.create_user(
            "c_lead", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="C POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        POCMembership.objects.create(poc=self.poc, user=self.lead, role_in_poc="lead")
        self.uc = UseCase.objects.create(poc=self.poc, title="UC", created_by=self.admin)
        self.req = Requirement.objects.create(
            poc=self.poc, req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )

    def test_member_can_comment_and_halo_appears(self):
        self.assertFalse(self.uc.has_open_comment)
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:comment_create", args=["usecase", self.uc.pk]),
            {"text": "Please clarify the actor."},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(self.uc.has_open_comment)

    def test_non_member_cannot_comment(self):
        outsider = User.objects.create_user(
            "c_outsider", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.client.force_login(outsider)
        resp = self.client.post(
            reverse("pocs:comment_create", args=["requirement", self.req.pk]),
            {"text": "Trying to comment."},
        )
        self.assertEqual(resp.status_code, 403)
        self.assertFalse(self.req.has_open_comment)

    def test_lead_can_close_thread_clears_halo(self):
        from django.contrib.contenttypes.models import ContentType

        from apps.pocs.models import Comment

        comment = Comment.objects.create(
            content_type=ContentType.objects.get_for_model(Requirement),
            object_id=self.req.pk,
            poc=self.poc,
            text="Missing detail.",
            author=self.member,
        )
        self.assertTrue(self.req.has_open_comment)
        self.client.force_login(self.lead)
        resp = self.client.post(reverse("pocs:comment_close", args=[comment.pk]))
        self.assertEqual(resp.status_code, 200)
        comment.refresh_from_db()
        self.assertEqual(comment.status, "closed")
        self.assertEqual(comment.closed_by, self.lead)
        self.assertIsNotNone(comment.closed_at)
        self.assertFalse(self.req.has_open_comment)

    def test_member_cannot_close_thread(self):
        from apps.pocs.models import Comment
        from django.contrib.contenttypes.models import ContentType

        comment = Comment.objects.create(
            content_type=ContentType.objects.get_for_model(UseCase),
            object_id=self.uc.pk,
            poc=self.poc,
            text="Needs review.",
            author=self.member,
        )
        self.client.force_login(self.member)
        resp = self.client.post(reverse("pocs:comment_close", args=[comment.pk]))
        self.assertEqual(resp.status_code, 403)
        comment.refresh_from_db()
        self.assertEqual(comment.status, "open")

    def test_any_member_can_reply_to_open_thread(self):
        from django.contrib.contenttypes.models import ContentType

        from apps.pocs.models import Comment

        comment = Comment.objects.create(
            content_type=ContentType.objects.get_for_model(Requirement),
            object_id=self.req.pk,
            poc=self.poc,
            text="Missing detail.",
            author=self.lead,
        )
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:comment_reply", args=[comment.pk]), {"text": "On it."}
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(comment.messages.count(), 1)
        self.assertEqual(comment.messages.first().author, self.member)

    def test_cannot_reply_to_closed_thread(self):
        from django.contrib.contenttypes.models import ContentType

        from apps.pocs.models import Comment

        comment = Comment.objects.create(
            content_type=ContentType.objects.get_for_model(Requirement),
            object_id=self.req.pk,
            poc=self.poc,
            text="Missing detail.",
            author=self.member,
        )
        comment.close(self.lead)
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:comment_reply", args=[comment.pk]), {"text": "Too late."}
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(comment.messages.count(), 0)

    def test_central_page_lists_all_threads_grouped_by_poc_and_tab(self):
        from django.contrib.contenttypes.models import ContentType

        from apps.pocs.models import Comment, Test

        Comment.objects.create(
            content_type=ContentType.objects.get_for_model(Requirement),
            object_id=self.req.pk, poc=self.poc, text="Spec thread", author=self.member,
        )
        phase = Phase.objects.create(poc=self.poc, name="System Testing", order=1)
        test = Test.objects.create(phase=phase, title="A test")
        Comment.objects.create(
            content_type=ContentType.objects.get_for_model(Test),
            object_id=test.pk, poc=self.poc, text="Test thread", author=self.member,
        )
        # A plain member (not a lead/admin) can still browse everything.
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:comments"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Spec thread")
        self.assertContains(resp, "Test thread")
        self.assertContains(resp, self.poc.name)


class EntityPreviewAndQuickStatusTests(TestCase):
    """Section 0 (shared preview modal) + Section 1 (quick UseCase status
    change from the popup): cross-reference chips render as clickable
    previews, and usecase_set_status behaves like the full edit form."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "ep_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.lead = User.objects.create_user(
            "ep_lead", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.member = User.objects.create_user(
            "ep_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="EP POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.lead, role_in_poc="lead")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.req = Requirement.objects.create(
            poc=self.poc, req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )
        self.uc = UseCase.objects.create(poc=self.poc, title="UC", created_by=self.admin)
        self.uc.requirements.add(self.req)

    def test_usecase_preview_links_requirement_as_clickable_chip(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:usecase_preview", args=[self.uc.pk]))
        self.assertContains(resp, "openEntityPreview(")
        # Chip shows req_id (spec item 9), not the audit-trail code.
        self.assertContains(resp, self.req.req_id)

    def test_requirement_preview_links_usecase_as_clickable_chip(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:requirement_preview", args=[self.req.pk]))
        self.assertContains(resp, "openEntityPreview(")
        self.assertContains(resp, self.uc.code)

    def test_member_preview_has_no_status_editor(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:usecase_preview", args=[self.uc.pk]))
        self.assertNotContains(resp, "usecase_set_status")

    def test_lead_preview_has_status_editor(self):
        self.client.force_login(self.lead)
        resp = self.client.get(reverse("pocs:usecase_preview", args=[self.uc.pk]))
        self.assertContains(resp, reverse("pocs:usecase_set_status", args=[self.uc.pk]))

    def test_lead_can_quick_change_status(self):
        from apps.pocs.models import AuditLog

        self.client.force_login(self.lead)
        resp = self.client.post(
            reverse("pocs:usecase_set_status", args=[self.uc.pk]),
            {"status": "active"},
        )
        self.assertEqual(resp.status_code, 200)
        self.uc.refresh_from_db()
        self.assertEqual(self.uc.status, "active")
        self.assertEqual(self.uc.modified_by, self.lead)
        self.assertTrue(
            AuditLog.objects.filter(poc=self.poc, action="usecase_updated").exists()
        )

    def test_member_cannot_quick_change_status(self):
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:usecase_set_status", args=[self.uc.pk]),
            {"status": "active"},
        )
        self.assertEqual(resp.status_code, 403)
        self.uc.refresh_from_db()
        self.assertEqual(self.uc.status, "draft")

    def test_quick_status_blocked_when_poc_closed(self):
        self.poc.closure_date = timezone.now().date()
        self.poc.closed_at = timezone.now()
        self.poc.closed_by = self.admin
        self.poc.save()
        self.client.force_login(self.lead)
        resp = self.client.post(
            reverse("pocs:usecase_set_status", args=[self.uc.pk]),
            {"status": "active"},
        )
        self.assertEqual(resp.status_code, 403)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class NonLeadReportGenerationTests(TestCase):
    """Section 2 fix: a non-lead POC member must be able to generate/download
    the SAME full Functional Analysis report a lead would get — the
    "Generate report" action was previously hidden from them (can_edit-gated
    in phase_detail.html) even though the view/generation code itself never
    filtered content by role, so download for a non-lead came back empty
    (nothing had ever been generated for them to download)."""

    def setUp(self):
        from apps.reports.models import ReportSettings

        self.admin = User.objects.create_user(
            "nl_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "nl_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="NL POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(
            poc=self.poc, name="FA", order=1, kind=PhaseKind.FUNCTIONAL_ANALYSIS
        )
        FunctionalAnalysisStep.objects.create(title="Step One", order=1)
        settings_row = ReportSettings.load()
        settings_row.default_template.save(
            "default.docx", SimpleUploadedFile("default.docx", _minimal_docx_bytes()), save=True
        )

    def test_non_lead_member_sees_generate_button(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Generate report")

    def test_non_lead_member_generates_full_report(self):
        from docx import Document

        from apps.reports.models import GeneratedReport

        self.client.force_login(self.member)
        resp = self.client.post(reverse("reports:phase_documents", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 302)

        report = GeneratedReport.objects.filter(phase=self.phase).latest("requested_at")
        self.assertEqual(report.status, GeneratedReport.Status.READY, report.error_message)
        self.assertTrue(report.output_file)

        report.output_file.open("rb")
        try:
            doc = Document(BytesIO(report.output_file.read()))
        finally:
            report.output_file.close()
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Step One", full_text)

        # Downloading it back as the same non-lead member works and is not empty.
        dl = self.client.get(reverse("reports:download", args=[report.pk]))
        self.assertEqual(dl.status_code, 200)
        self.assertGreater(len(b"".join(dl.streaming_content)), 0)


_MINIMAL_SVG = b"""<svg xmlns="http://www.w3.org/2000/svg" width="120" height="60">
<rect x="5" y="5" width="50" height="30" fill="green"/>
<text x="10" y="50" font-size="10">chart</text>
</svg>"""


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class SvgReportEmbedTests(TestCase):
    """Section 4: SVG can be attached/embedded like any other image, and the
    generated .docx converts it to a raster image (Word has no native SVG
    support) instead of silently dropping it or breaking generation."""

    def setUp(self):
        from apps.reports.models import ReportSettings

        self.admin = User.objects.create_user(
            "svg_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="SVG POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(
            poc=self.poc, name="Docs", order=1, kind=PhaseKind.DOCUMENTATION
        )
        settings_row = ReportSettings.load()
        settings_row.default_template.save(
            "default.docx", SimpleUploadedFile("default.docx", _minimal_docx_bytes()), save=True
        )
        self.image = PhaseImage.objects.create(
            phase=self.phase,
            image=SimpleUploadedFile("chart.svg", _MINIMAL_SVG, content_type="image/svg+xml"),
        )
        PhaseDocument.objects.create(
            phase=self.phase,
            title="Section 1",
            content=f"Some text\n\n{self.image.markdown_snippet}\n\nMore text",
            order=1,
        )

    def test_svg_upload_passes_model_validation(self):
        self.image.full_clean()  # raises ValidationError if .svg were rejected

    def test_report_rasterizes_svg_to_embedded_image(self):
        from docx import Document

        from apps.reports.generation import generate_phase_report_from_documents
        from apps.reports.models import GeneratedReport

        report = generate_phase_report_from_documents(self.phase, self.admin)
        self.assertEqual(
            report.status, GeneratedReport.Status.READY, report.error_message
        )
        report.output_file.open("rb")
        try:
            doc = Document(BytesIO(report.output_file.read()))
        finally:
            report.output_file.close()
        # Converted (not dropped, not left as literal "[image: ...]" alt text).
        self.assertGreaterEqual(len(doc.inline_shapes), 1)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("[image:", full_text)

    def test_broken_svg_falls_back_to_alt_text_without_crashing(self):
        from apps.reports.converter.builder import build_docx

        blocks = [
            {"type": "paragraph", "children": [{"text": "before"}]},
            {"type": "image", "url": "/media/phase_images/broken.svg", "alt": "broken chart"},
            {"type": "paragraph", "children": [{"text": "after"}]},
        ]
        broken_path = os.path.join(tempfile.mkdtemp(), "broken.svg")
        with open(broken_path, "wb") as f:
            f.write(b"not actually an svg file")

        out = build_docx(blocks, image_resolver=lambda url: broken_path)
        self.assertTrue(out)  # generation completed, didn't raise


class TestReorderTests(TestCase):
    """Improvement brief item 10: drag-and-drop reordering of a phase's tests
    also renumbers test_code to match, so ST-NNN tracks display/execution
    order rather than creation order."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "reorder_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "reorder_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="Reorder POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="System Testing", order=1)
        self.t1 = Test.objects.create(phase=self.phase, title="First")
        self.t2 = Test.objects.create(phase=self.phase, title="Second")
        self.t3 = Test.objects.create(phase=self.phase, title="Third")

    def test_tests_get_sequential_codes_on_creation(self):
        self.assertEqual(self.t1.test_code, "ST-001")
        self.assertEqual(self.t2.test_code, "ST-002")
        self.assertEqual(self.t3.test_code, "ST-003")
        # Default ordering sorts by test_code (spec item 10 follow-up), so
        # display always reads ST-001, ST-002, … regardless of order/id.
        self.assertEqual(list(self.phase.tests.all()), [self.t1, self.t2, self.t3])

    def test_reorder_renumbers_codes_and_persists_new_order(self):
        self.client.force_login(self.admin)
        resp = self.client.post(
            reverse("pocs:test_reorder", args=[self.phase.pk]),
            {"test_ids": [self.t3.pk, self.t1.pk, self.t2.pk]},
        )
        self.assertEqual(resp.status_code, 200)
        self.t1.refresh_from_db()
        self.t2.refresh_from_db()
        self.t3.refresh_from_db()
        # Same set of numeric suffixes, reassigned to the new positions.
        self.assertEqual(self.t3.test_code, "ST-001")
        self.assertEqual(self.t1.test_code, "ST-002")
        self.assertEqual(self.t2.test_code, "ST-003")
        self.assertEqual([self.t3.order, self.t1.order, self.t2.order], [0, 1, 2])
        # The phase's default ordering now reflects the drop order.
        self.assertEqual(list(self.phase.tests.all()), [self.t3, self.t1, self.t2])

    def test_out_of_order_ids_still_sort_by_code_by_default(self):
        # Reproduces the reported bug: even if id/creation order and the
        # eventual test_code assignment ever drift apart, the default
        # queryset must still read ST-001, ST-002, ST-003 — not id order.
        Test.objects.filter(pk=self.t1.pk).update(test_code="ST-003")
        Test.objects.filter(pk=self.t3.pk).update(test_code="ST-001")
        self.assertEqual(
            [t.pk for t in self.phase.tests.all()], [self.t3.pk, self.t2.pk, self.t1.pk]
        )

    def test_define_test_order_button_and_modal_render_for_lead(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Define test order")
        self.assertContains(resp, 'id="test-order-list"')
        self.assertContains(resp, self.t1.title)
        self.assertContains(resp, self.t1.test_code)

    def test_define_test_order_button_hidden_for_plain_member(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, "Define test order")

    def test_member_cannot_reorder(self):
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:test_reorder", args=[self.phase.pk]),
            {"test_ids": [self.t3.pk, self.t1.pk, self.t2.pk]},
        )
        self.assertEqual(resp.status_code, 403)


class TaskGanttRenderTests(TestCase):
    """Section 5: phase_detail and the global Tasks page render the tree+Gantt
    without duplicating a sub-task (the bug being fixed — a sub-task used to
    also appear as its own top-level entry on the global Tasks page)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "gr_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "gr_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="GR POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        self.parent = Task.objects.create(phase=self.phase, title="Parent Task Alpha")
        self.sub = Task.objects.create(
            phase=self.phase, title="Sub Task Bravo", parent=self.parent
        )

    def test_phase_detail_no_longer_renders_tasks(self):
        # Tasks are removed from the phase page (spec item 4 — minimalism);
        # the duplicate-render regression guard now lives on the Tasks
        # workspace (see test_global_tasks_page_shows_hierarchy_not_duplicate).
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode()
        self.assertNotIn(f'id="task-{self.sub.pk}"', body)
        self.assertNotIn(f'id="task-{self.parent.pk}"', body)

    def test_global_tasks_page_shows_hierarchy_not_duplicate(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:tasks"))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode()
        self.assertEqual(body.count(f'id="task-{self.sub.pk}"'), 1)
        self.assertEqual(body.count(f'id="task-{self.parent.pk}"'), 1)

    def test_phase_card_shows_own_task_summary(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:detail", args=[self.poc.pk]) + "?tab=phases")
        self.assertContains(resp, "0/2 tasks")


class DarkModeSmokeTests(TestCase):
    """Section 6: theme toggle + semantic tokens are present and wired up on
    both the authenticated app shell and the (separately-headed) login page."""

    def test_login_page_has_theme_toggle_and_tokens(self):
        resp = self.client.get(reverse("accounts:login"))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode()
        self.assertIn("theme-toggle", body)
        self.assertIn("html.dark", body)
        self.assertIn("prefers-color-scheme", body)

    def test_app_shell_has_theme_toggle_and_tokens(self):
        admin = User.objects.create_user(
            "dm_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.client.force_login(admin)
        resp = self.client.get(reverse("core:dashboard"))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode()
        self.assertIn("html.dark", body)
        self.assertIn("localStorage.setItem('theme'", body)


class UnsavedChangesGuardTests(TestCase):
    """Improvement brief item 7: a beforeunload/internal-link guard on the
    app's edit forms — smoke-checks the wiring is present, not the JS
    behaviour itself (unreachable from a Django TestCase)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "ug_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="Guard POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(poc=self.poc, name="P", order=1)
        self.client.force_login(self.admin)

    def test_guard_script_loaded_on_app_shell(self):
        resp = self.client.get(reverse("core:dashboard"))
        self.assertContains(resp, "unsaved_changes_guard.js")

    def test_poc_edit_form_has_guard_attribute(self):
        resp = self.client.get(reverse("pocs:edit", args=[self.poc.pk]))
        self.assertContains(resp, "data-unsaved-guard")

    def test_phase_edit_form_has_guard_attribute(self):
        resp = self.client.get(reverse("pocs:phase_edit", args=[self.phase.pk]))
        self.assertContains(resp, "data-unsaved-guard")

    def test_requirement_edit_form_has_guard_attribute(self):
        req = Requirement.objects.create(
            poc=self.poc, req_gravity="imposes_mvp", req_operation="navigation",
            req_functional="performance", req_category="normal_operation",
            created_by=self.admin,
        )
        resp = self.client.get(reverse("pocs:requirement_edit", args=[req.pk]))
        self.assertContains(resp, "data-unsaved-guard")


class MarkdownExportImportTests(TestCase):
    """Improvement brief item 1: Requirements & Use Cases export/import also
    accepts .md, in addition to the existing .xlsx round-trip."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "md_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="MD POC", created_by=self.admin, status="active")
        self.req = Requirement.objects.create(
            poc=self.poc, sub_system="SCADA", req_gravity="imposes_mvp",
            req_operation="navigation", req_functional="performance",
            req_category="normal_operation", description="Has a | pipe\nand a newline",
            created_by=self.admin,
        )
        self.uc = UseCase.objects.create(poc=self.poc, title="Login", created_by=self.admin)

    def test_requirement_export_md_round_trips_through_parser(self):
        from apps.pocs.requirements_import import (
            build_requirements_export_md,
            parse_requirements_md,
        )

        md = build_requirements_export_md(self.poc)
        self.assertIn("| code |", md)
        self.assertIn(self.req.code, md)
        # The pipe/newline in the description survive the escape/unescape round-trip.
        preview = parse_requirements_md(md, self.poc)
        self.assertEqual(len(preview), 1)
        row = preview[0]
        self.assertTrue(row["valid"], row["errors"])
        self.assertEqual(row["data"]["description"], "Has a | pipe\nand a newline")
        self.assertEqual(row["data"]["existing_id"], self.req.id)

    def test_usecase_export_md_round_trips_through_parser(self):
        from apps.pocs.requirements_import import (
            build_usecases_export_md,
            parse_usecases_md,
        )

        md = build_usecases_export_md(self.poc)
        self.assertIn(self.uc.code, md)
        preview = parse_usecases_md(md, self.poc)
        self.assertEqual(len(preview), 1)
        self.assertTrue(preview[0]["valid"], preview[0]["errors"])
        self.assertEqual(preview[0]["data"]["existing_id"], self.uc.id)

    def test_requirement_export_view_offers_md_format(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:requirement_export", args=[self.poc.pk]) + "?format=md")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "text/markdown")
        self.assertIn(self.req.code, resp.content.decode())

    def test_usecase_export_view_offers_md_format(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:usecase_export", args=[self.poc.pk]) + "?format=md")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "text/markdown")
        self.assertIn(self.uc.code, resp.content.decode())

    def test_requirement_import_view_accepts_md_upload(self):
        from apps.pocs.requirements_import import build_requirements_export_md

        # Re-uploading this POC's own export (unedited) round-trips as an
        # UPDATE of the same requirement (matched by its ``code`` column) —
        # proving .md upload works end-to-end through the view, same as .xlsx.
        md = build_requirements_export_md(self.poc)
        upload = SimpleUploadedFile("reqs.md", md.encode("utf-8"), content_type="text/markdown")
        self.client.force_login(self.admin)
        resp = self.client.post(reverse("pocs:requirement_import", args=[self.poc.pk]), {"file": upload})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "1 valid")
        resp = self.client.post(reverse("pocs:requirement_import", args=[self.poc.pk]), {"confirm": "1"})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(self.poc.requirements.count(), 1)
        self.req.refresh_from_db()
        self.assertEqual(self.req.description, "Has a | pipe\nand a newline")


class PhaseExternalTests(TestCase):
    """Improvement brief item 4: 'External' phase status, coexisting with
    Not Applicable, with a reusable Team catalog."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "ext_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="Ext POC", created_by=self.admin, status="active")
        self.phase = Phase.objects.create(poc=self.poc, name="Site Survey", order=1)

    def test_mark_external_sets_flag_and_teams_without_touching_status(self):
        # Spec item 1: External is an independent flag — progress status
        # (default Pending here) is left untouched.
        vendor = Team.objects.create(name="Vendor QA")
        self.phase.mark_external(self.admin, [vendor])
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, Phase.Status.PENDING)
        self.assertTrue(self.phase.is_external)
        self.assertIn(vendor, self.phase.teams.all())
        self.assertIsNotNone(self.phase.external_at)
        self.assertEqual(self.phase.external_by, self.admin)

    def test_mark_external_clears_na_and_vice_versa(self):
        # External then NA — NA wins, external flag/fields are cleared.
        vendor = Team.objects.create(name="Vendor QA")
        self.phase.mark_external(self.admin, [vendor])
        self.phase.mark_na(self.admin, "Site inaccessible")
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, Phase.Status.NOT_APPLICABLE)
        self.assertFalse(self.phase.is_external)
        self.assertIsNone(self.phase.external_at)
        self.assertEqual(self.phase.teams.count(), 0)
        # NA then External — External wins, status reverts to Pending, na
        # fields are cleared.
        self.phase.mark_external(self.admin, [vendor])
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, Phase.Status.PENDING)
        self.assertTrue(self.phase.is_external)
        self.assertEqual(self.phase.na_reason, "")
        self.assertIsNone(self.phase.na_at)

    def test_unmark_external_reverts_flag_only(self):
        vendor = Team.objects.create(name="Vendor QA")
        self.phase.mark_external(self.admin, [vendor])
        self.phase.unmark_external()
        self.phase.refresh_from_db()
        self.assertFalse(self.phase.is_external)
        self.assertEqual(self.phase.status, Phase.Status.PENDING)
        self.assertEqual(self.phase.teams.count(), 0)

    def test_external_leaf_uses_real_status_for_parent_rollup(self):
        # Spec item 1: marking a child External does NOT automatically count
        # it as done for the parent rollup — its real progress status does.
        parent = Phase.objects.create(poc=self.poc, name="Parent", order=0)
        self.phase.parent = parent
        self.phase.save()
        vendor = Team.objects.create(name="Vendor QA")
        self.phase.mark_external(self.admin, [vendor])
        parent.refresh_from_db()
        self.assertEqual(parent.compute_status(), Phase.Status.PENDING)

        # Progress on the External phase is still independently editable —
        # once its own status is set to Completed, the parent reflects it.
        self.phase.set_status_cascade(Phase.Status.COMPLETED)
        parent.refresh_from_db()
        self.assertEqual(parent.compute_status(), Phase.Status.COMPLETED)

    def test_mark_external_view_creates_new_team_inline(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:phase_mark_external", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 200)
        resp = self.client.post(
            reverse("pocs:phase_mark_external", args=[self.phase.pk]),
            {"new_teams": "Vendor QA, Site Ops"},
        )
        self.assertEqual(resp.status_code, 302)
        self.phase.refresh_from_db()
        self.assertTrue(self.phase.is_external)
        self.assertEqual(
            set(self.phase.teams.values_list("name", flat=True)), {"Vendor QA", "Site Ops"}
        )
        self.assertEqual(Team.objects.count(), 2)

    def test_mark_external_view_requires_at_least_one_team(self):
        self.client.force_login(self.admin)
        resp = self.client.post(reverse("pocs:phase_mark_external", args=[self.phase.pk]), {})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Select or add at least one team")

    def test_unmark_external_view(self):
        vendor = Team.objects.create(name="Vendor QA")
        self.phase.mark_external(self.admin, [vendor])
        self.client.force_login(self.admin)
        resp = self.client.post(reverse("pocs:phase_unmark_external", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 302)
        self.phase.refresh_from_db()
        self.assertEqual(self.phase.status, Phase.Status.PENDING)

    def test_member_cannot_mark_external(self):
        member = User.objects.create_user(
            "ext_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        POCMembership.objects.create(poc=self.poc, user=member, role_in_poc="member")
        self.client.force_login(member)
        resp = self.client.get(reverse("pocs:phase_mark_external", args=[self.phase.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_phase_detail_shows_external_banner_and_teams(self):
        vendor = Team.objects.create(name="Vendor QA")
        self.phase.mark_external(self.admin, [vendor])
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:phase_detail", args=[self.phase.pk]))
        self.assertContains(resp, "Marked")
        self.assertContains(resp, "External")
        self.assertContains(resp, "Vendor QA")

    def test_final_report_shows_external_team(self):
        from apps.reports.generation import build_final_report_markdown

        vendor = Team.objects.create(name="Vendor QA")
        self.phase.mark_external(self.admin, [vendor])
        md = build_final_report_markdown(self.poc, "done")
        self.assertIn("External — Vendor QA", md)


class UCRequirementMatrixTests(TestCase):
    """Improvement brief item 3: UC×Requirement matrix (+ item 7's shared
    hover-card, reused here on the row/column headers)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "mx_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.member = User.objects.create_user(
            "mx_member", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.poc = POC.objects.create(name="Matrix POC", created_by=self.admin, status="active")
        POCMembership.objects.create(poc=self.poc, user=self.member, role_in_poc="member")
        self.uc = UseCase.objects.create(
            poc=self.poc, title="Operator logs in", description="Full UC description",
            created_by=self.admin,
        )
        self.req1 = Requirement.objects.create(
            poc=self.poc, sub_system="SCADA", req_gravity="imposes_mvp",
            req_operation="navigation", req_functional="performance",
            req_category="normal_operation", description="Req 1 description",
            created_by=self.admin,
        )
        self.req2 = Requirement.objects.create(
            poc=self.poc, sub_system="HMI", req_gravity="imposes_mvp",
            req_operation="cybersecurity", req_functional="performance",
            req_category="degraded_operation", created_by=self.admin,
        )

    def test_matrix_renders_rows_cols_and_hover_card_data(self):
        self.client.force_login(self.member)
        resp = self.client.get(reverse("pocs:uc_requirement_matrix", args=[self.poc.pk]))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode()
        self.assertIn(self.uc.code, body)
        # Matrix columns show req_id (spec item 9), not the audit-trail code.
        self.assertIn(self.req1.req_id, body)
        self.assertIn(self.req2.req_id, body)
        self.assertIn("data-hover-card", body)
        self.assertIn("Full UC description", body)
        self.assertIn("Req 1 description", body)
        # Different categories get different colour tokens (grouped/coloured, spec item 3).
        self.assertNotEqual(
            self.req1.req_category, self.req2.req_category
        )

    def test_member_cannot_toggle_link(self):
        self.client.force_login(self.member)
        resp = self.client.post(
            reverse("pocs:matrix_toggle_link", args=[self.poc.pk]),
            {"usecase_id": self.uc.pk, "requirement_id": self.req1.pk},
        )
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(self.uc.requirements.count(), 0)

    def test_lead_can_toggle_link_on_and_off(self):
        self.client.force_login(self.admin)
        url = reverse("pocs:matrix_toggle_link", args=[self.poc.pk])
        resp = self.client.post(url, {"usecase_id": self.uc.pk, "requirement_id": self.req1.pk})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json(), {"linked": True})
        self.assertIn(self.req1, self.uc.requirements.all())

        resp = self.client.post(url, {"usecase_id": self.uc.pk, "requirement_id": self.req1.pk})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json(), {"linked": False})
        self.assertNotIn(self.req1, self.uc.requirements.all())

    def test_matrix_reflects_existing_links(self):
        self.uc.requirements.add(self.req1)
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:uc_requirement_matrix", args=[self.poc.pk]))
        body = resp.content.decode()
        # The linked cell renders a filled check icon.
        self.assertIn('data-lucide="check"', body)


class EntityHoverPreviewBriefTests(TestCase):
    """A follow-up request: the delayed hover-card (used on the graph and the
    matrix) should also work on every chip that links a Requirement/Use Case/
    Test elsewhere in the app. Implemented as a `?brief=1` JSON fast-path on
    the existing preview endpoints (apps/pocs/views.py), fetched generically
    by static/js/entity_hover_preview.js — these tests cover the server side."""

    def setUp(self):
        self.admin = User.objects.create_user(
            "hp_admin", password="x", is_active=True, role=User.Role.ADMIN
        )
        self.poc = POC.objects.create(name="Hover POC", created_by=self.admin, status="active")
        self.uc = UseCase.objects.create(
            poc=self.poc, title="Login", description="UC full description", created_by=self.admin,
        )
        self.req = Requirement.objects.create(
            poc=self.poc, sub_system="SCADA", req_gravity="imposes_mvp",
            req_operation="navigation", req_functional="performance",
            req_category="normal_operation", description="Req full description",
            created_by=self.admin,
        )
        self.phase = Phase.objects.create(poc=self.poc, name="Testing", order=1)
        self.test = Test.objects.create(
            phase=self.phase, title="Smoke test", test_code="UT-001",
            description="Test full description",
        )

    def test_requirement_preview_brief_returns_json(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:requirement_preview", args=[self.req.pk]) + "?brief=1")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/json")
        # Hover-card title is req_id (spec item 9), not the audit-trail code.
        self.assertEqual(resp.json(), {"title": self.req.req_id, "description": "Req full description"})

    def test_usecase_preview_brief_returns_json(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:usecase_preview", args=[self.uc.pk]) + "?brief=1")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json(), {"title": f"{self.uc.code} · {self.uc.title}", "description": "UC full description"})

    def test_test_preview_brief_returns_json(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:test_preview", args=[self.test.pk]) + "?brief=1")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json(), {"title": "UT-001", "description": "Test full description"})

    def test_brief_still_requires_poc_membership(self):
        outsider = User.objects.create_user(
            "hp_outsider", password="x", is_active=True, role=User.Role.TEAM_MEMBER
        )
        self.client.force_login(outsider)
        resp = self.client.get(reverse("pocs:requirement_preview", args=[self.req.pk]) + "?brief=1")
        self.assertEqual(resp.status_code, 403)

    def test_normal_preview_still_returns_html_without_brief_param(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("pocs:requirement_preview", args=[self.req.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "text/html; charset=utf-8")
