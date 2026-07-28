"""
Requirement import from Excel (spec Fase 10c).

``parse_requirements_xlsx`` reads an .xlsx into a per-row preview with row-level
validation (which field fails and why); ``import_requirements`` then creates the
valid rows. The ``code`` is auto-generated (as in the UI) and the POC/user come
from the view — the sheet only carries the descriptive fields of Fase 2.

The classification columns (gravity/operation/functional/category) are
POC-extensible: a value that isn't one of the built-in suggestions is accepted
as-is rather than rejected, and registered as a :class:`RequirementFieldOption`
so it becomes a suggestion on the manual form too (see ``RequirementForm``'s
"+ Add new…" control).
"""

from .audit import record_audit
from .models import Requirement, RequirementFieldOption, UseCase

_CLASSIFICATION_FIELDS = ("req_gravity", "req_operation", "req_functional", "req_category")
_OPTIONAL_CLASSIFICATION_FIELDS = ("life_cycle_phase",)


def _reverse_map(choices):
    """Accept either the stored value or the human label (case-insensitive)."""
    m = {}
    for value, label in choices:
        m[value.lower()] = value
        m[label.lower()] = value
    return m


_GRAVITY = _reverse_map(Requirement.Gravity.choices)
_OPERATION = _reverse_map(Requirement.Operation.choices)
_FUNCTIONAL = _reverse_map(Requirement.Functional.choices)
_CATEGORY = _reverse_map(Requirement.Category.choices)
_LIFE_CYCLE_PHASE = _reverse_map(Requirement.LifeCyclePhase.choices)

_PREDEFINED_MAPS = {
    "req_gravity": _GRAVITY,
    "req_operation": _OPERATION,
    "req_functional": _FUNCTIONAL,
    "req_category": _CATEGORY,
    "life_cycle_phase": _LIFE_CYCLE_PHASE,
}


def _custom_map(poc, field_name):
    """This POC's previously-added custom values for ``field_name``, by lowercase."""
    if poc is None:
        return {}
    return {
        v.lower(): v
        for v in poc.requirement_field_options.filter(field=field_name).values_list(
            "value", flat=True
        )
    }


# Header name → model field (all matched case-insensitively).
_HEADER_ALIASES = {
    # ``code`` is Light's own auto-generated identifier — present only on an
    # exported sheet (see ``build_requirements_export_xlsx``). When a row's
    # code matches an existing requirement of this POC, the row UPDATES it
    # instead of creating a new one (the round-trip export→edit→import flow).
    "code": {"code"},
    "external_code": {"external_code", "external code", "source code", "original code"},
    "sub_system": {"sub_system", "subsystem", "sub system", "system"},
    "req_gravity": {"req_gravity", "gravity"},
    "req_operation": {"req_operation", "operation"},
    "req_functional": {"req_functional", "functional"},
    "req_category": {"req_category", "category"},
    "description": {"description", "desc"},
    "validation_criteria": {"validation_criteria", "validation criteria", "validation"},
    "life_cycle_phase": {"life_cycle_phase", "life cycle phase", "lifecycle"},
    "reference_documentations": {"reference_documentations", "references", "reference"},
    "remarks": {"remarks", "notes"},
    "use_cases": {"use_cases", "use cases", "usecases", "use_case_codes"},
}


def _norm(value):
    return str(value if value is not None else "").strip()


# ---------------------------------------------------------------------------
# Markdown table read/write — the .md counterpart of the .xlsx round-trip.
# A GitHub-flavoured pipe table with the exact same header names as the .xlsx
# format, so the two are interchangeable for both export and import.
# ---------------------------------------------------------------------------
def _md_escape(value):
    # An HTML entity (not a backslash) so a naive split("|") below never
    # re-breaks an escaped pipe back into two cells.
    return str(value if value is not None else "").replace("|", "&#124;").replace("\n", "<br>")


def _md_unescape(value):
    return value.replace("<br>", "\n").replace("&#124;", "|")


def _split_md_row(line):
    cells = line.strip()
    if cells.startswith("|"):
        cells = cells[1:]
    if cells.endswith("|"):
        cells = cells[:-1]
    return [_md_unescape(c.strip()) for c in cells.split("|")]


def _read_md_table(file):
    """Parse a single pipe-table out of a Markdown file/text: (headers, rows).

    Only the first table found is read (matches what ``build_*_export_md``
    writes) — any other content on the page is ignored.
    """
    text = file.read() if hasattr(file, "read") else file
    if isinstance(text, bytes):
        text = text.decode("utf-8-sig")
    lines = [ln for ln in text.splitlines() if ln.strip().startswith("|")]
    if len(lines) < 2:
        return [], []
    headers = _split_md_row(lines[0])
    # lines[1] is the "| --- | --- |" separator — skip it.
    return headers, [_split_md_row(ln) for ln in lines[2:]]


def _write_md_table(headers, rows):
    lines = [
        "| " + " | ".join(_md_escape(h) for h in headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(_md_escape(c) for c in row) + " |")
    return "\n".join(lines) + "\n"


def _rows_from_xlsx(file):
    """(headers, data_rows) from an .xlsx — shared by the requirement/use
    case parsers below."""
    import openpyxl

    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = next(rows_iter, None)
    if not headers:
        return [], []
    data_rows = [
        row for row in rows_iter
        if row is not None and not all(_norm(c) == "" for c in row)
    ]
    return list(headers), data_rows


def parse_requirements_xlsx(file, poc=None):
    """Return a list of preview rows: {row, data, errors, valid} from an .xlsx."""
    headers, data_rows = _rows_from_xlsx(file)
    if not headers:
        return []
    return _requirement_preview_rows(headers, data_rows, poc)


def parse_requirements_md(file, poc=None):
    """Same as ``parse_requirements_xlsx`` but reading a Markdown pipe table
    (as written by ``build_requirements_export_md``)."""
    headers, data_rows = _read_md_table(file)
    if not headers:
        return []
    return _requirement_preview_rows(headers, data_rows, poc)


def _requirement_preview_rows(headers, data_rows, poc=None):
    """Return a list of preview rows: {row, data, errors, valid}.

    ``poc`` supplies the POC's own previously-added custom classification
    values, so a value the POC has used before resolves to its existing exact
    spelling instead of creating a near-duplicate. A value seen for the first
    time is accepted as-is (see module docstring).

    An optional ``use_cases`` column may list existing Use Case codes of this
    POC (comma-separated) to link on import — left out or blank, a requirement
    is imported with no use cases linked (the previous, only behaviour).
    """
    col_field = {}
    for idx, h in enumerate(headers):
        hn = _norm(h).lower()
        for field, aliases in _HEADER_ALIASES.items():
            if hn in aliases:
                col_field[idx] = field
    has_uc_column = "use_cases" in col_field.values()

    custom_maps = {
        f: _custom_map(poc, f) for f in (*_CLASSIFICATION_FIELDS, *_OPTIONAL_CLASSIFICATION_FIELDS)
    }
    uc_by_code = {uc.code: uc.id for uc in poc.use_cases.all()} if poc is not None else {}
    req_by_code = {}
    if poc is not None:
        req_by_code = dict(poc.requirements.values_list("code", "id"))

    # A requirement has no natural business key, but two rows sharing the same
    # validation criteria are, in practice, the same requirement — used to
    # reject re-imports of the same sheet (or accidental copies) as duplicates.
    # Mapped to the owning requirement's id so an UPDATE row (matched by
    # ``code``) doesn't get flagged as a duplicate of itself.
    existing_criteria = {}
    if poc is not None:
        existing_criteria = dict(
            poc.requirements.exclude(validation_criteria="")
            .values_list("validation_criteria", "id")
        )
    seen_criteria = set()

    preview = []
    for r_i, row in enumerate(data_rows, start=2):
        if row is None or all(_norm(c) == "" for c in row):
            continue
        data = {}
        for idx, field in col_field.items():
            if idx < len(row):
                data[field] = _norm(row[idx])

        errors = []

        # A ``code`` matching an existing requirement of this POC turns the
        # row into an UPDATE of that requirement instead of a new one.
        existing_id = None
        code = data.get("code", "")
        if code:
            existing_id = req_by_code.get(code)
            if existing_id is None:
                errors.append(f"Unknown requirement code: {code}")

        def resolve(field, label, required=True):
            raw = data.get(field, "")
            if not raw:
                if required:
                    errors.append(f"{label} is required")
                return ""
            key = raw.lower()
            if key in _PREDEFINED_MAPS[field]:
                return _PREDEFINED_MAPS[field][key]
            if key in custom_maps[field]:
                return custom_maps[field][key]
            # Not a known value yet — accepted as a new custom one (registered
            # on import so it becomes a suggestion from now on).
            return raw

        # ``use_case_ids`` stays ``None`` when the sheet has no ``use_cases``
        # column at all — an UPDATE row then leaves existing links untouched
        # instead of wiping them; a blank cell (column present, empty) means
        # "link none", same as before.
        use_case_ids, unknown_ucs = (None, []) if not has_uc_column else ([], [])
        if has_uc_column:
            for uc_code in (data.get("use_cases", "") or "").split(","):
                uc_code = uc_code.strip()
                if not uc_code:
                    continue
                if uc_code in uc_by_code:
                    use_case_ids.append(uc_by_code[uc_code])
                else:
                    unknown_ucs.append(uc_code)
        if unknown_ucs:
            errors.append("Unknown use case code(s): " + ", ".join(unknown_ucs))

        validation_criteria = data.get("validation_criteria", "")
        if validation_criteria:
            owner_id = existing_criteria.get(validation_criteria)
            if (owner_id is not None and owner_id != existing_id) or validation_criteria in seen_criteria:
                errors.append(
                    "Duplicate requirement — this validation criteria already exists"
                )
            else:
                seen_criteria.add(validation_criteria)

        clean = {
            "external_code": data.get("external_code", ""),
            "sub_system": data.get("sub_system", ""),
            "req_gravity": resolve("req_gravity", "Gravity"),
            "req_operation": resolve("req_operation", "Operation"),
            "req_functional": resolve("req_functional", "Functional"),
            "req_category": resolve("req_category", "Category"),
            "description": data.get("description", ""),
            "validation_criteria": validation_criteria,
            "life_cycle_phase": resolve("life_cycle_phase", "Lifecycle status", required=False),
            "reference_documentations": data.get("reference_documentations", ""),
            "remarks": data.get("remarks", ""),
            "use_case_ids": use_case_ids,
            "existing_id": existing_id,
        }
        preview.append(
            {"row": r_i, "data": clean, "errors": errors, "valid": not errors}
        )
    return preview


def import_requirements(poc, rows, user):
    """Create/update Requirements from validated preview rows.

    A row whose ``code`` matched an existing requirement of this POC (see
    ``parse_requirements_xlsx``) UPDATES it in place instead of creating a
    duplicate — this is what makes the export → edit → re-import round-trip
    work. Returns ``(created, updated)``.
    """
    created = updated = 0
    for row in rows:
        if not row.get("valid"):
            continue
        data = dict(row["data"])
        use_case_ids = data.pop("use_case_ids", [])
        existing_id = data.pop("existing_id", None)
        if existing_id:
            req = Requirement.objects.get(pk=existing_id)
            for field_name, value in data.items():
                setattr(req, field_name, value)
            req.modified_by = user
            req.save()
            updated += 1
            record_audit(req, "requirement_updated", user, {})
        else:
            req = Requirement.objects.create(
                poc=poc, created_by=user, modified_by=user, **data
            )
            created += 1
            record_audit(req, "requirement_created", user, {"code": {"before": None, "after": req.code}})
        if use_case_ids is not None:
            req.use_cases.set(use_case_ids)
        for field_name in (*_CLASSIFICATION_FIELDS, *_OPTIONAL_CLASSIFICATION_FIELDS):
            value = data.get(field_name)
            if value and value not in dict(Requirement.predefined_choices(field_name)):
                RequirementFieldOption.objects.get_or_create(
                    poc=poc, field=field_name, value=value
                )
    return created, updated


# ---------------------------------------------------------------------------
# Use Cases (spec Fase 10c)
# ---------------------------------------------------------------------------
_UC_PRIORITY = _reverse_map(UseCase.Priority.choices)
_UC_STATUS = _reverse_map(UseCase.Status.choices)

_UC_HEADER_ALIASES = {
    # ``code`` present only on an exported sheet — a matching row updates
    # that use case instead of creating a new one (see ``parse_requirements_xlsx``).
    "code": {"code"},
    "external_code": {"external_code", "external code", "source code", "original code"},
    "title": {"title", "name", "use case", "use_case"},
    "description": {"description", "desc"},
    "actor": {"actor", "role"},
    "priority": {"priority"},
    "status": {"status"},
    "remarks": {"remarks", "notes"},
    "requirements": {"requirements", "requirement", "reqs", "requirement_codes"},
}


def parse_usecases_xlsx(file, poc):
    """Return preview rows for a Use Case import (validated against ``poc``) from an .xlsx."""
    headers, data_rows = _rows_from_xlsx(file)
    if not headers:
        return []
    return _usecase_preview_rows(headers, data_rows, poc)


def parse_usecases_md(file, poc):
    """Same as ``parse_usecases_xlsx`` but reading a Markdown pipe table
    (as written by ``build_usecases_export_md``)."""
    headers, data_rows = _read_md_table(file)
    if not headers:
        return []
    return _usecase_preview_rows(headers, data_rows, poc)


def _usecase_preview_rows(headers, data_rows, poc):
    """Return preview rows for a Use Case import (validated against ``poc``).

    Title is required and must be unique within the POC (and the batch). Priority
    defaults to Medium and status to Draft when blank. A ``requirements`` column
    may list existing requirement codes (comma-separated) to link.
    """
    col_field = {}
    for idx, h in enumerate(headers):
        hn = _norm(h).lower()
        for field, aliases in _UC_HEADER_ALIASES.items():
            if hn in aliases:
                col_field[idx] = field
    has_req_column = "requirements" in col_field.values()

    existing_titles = dict(poc.use_cases.values_list("title", "id"))
    uc_by_code = dict(poc.use_cases.values_list("code", "id"))
    req_by_code = {r.code: r.id for r in poc.requirements.all()}
    seen_titles = set()

    preview = []
    for r_i, row in enumerate(data_rows, start=2):
        if row is None or all(_norm(c) == "" for c in row):
            continue
        data = {}
        for idx, field in col_field.items():
            if idx < len(row):
                data[field] = _norm(row[idx])

        errors = []

        existing_id = None
        code = data.get("code", "")
        if code:
            existing_id = uc_by_code.get(code)
            if existing_id is None:
                errors.append(f"Unknown use case code: {code}")

        title = data.get("title", "")
        if not title:
            errors.append("Title is required")
        else:
            owner_id = existing_titles.get(title)
            if (owner_id is not None and owner_id != existing_id) or title.lower() in seen_titles:
                errors.append(f"Title “{title}” already exists in this POC")
            else:
                seen_titles.add(title.lower())

        def resolve_optional(field, mapping, label, default):
            raw = data.get(field, "")
            if not raw:
                return default
            val = mapping.get(raw.lower())
            if not val:
                errors.append(f"{label} “{raw}” is not a valid option")
                return default
            return val

        priority = resolve_optional("priority", _UC_PRIORITY, "Priority", "medium")
        status = resolve_optional("status", _UC_STATUS, "Status", "draft")

        # ``requirement_ids`` stays ``None`` when the sheet has no
        # ``requirements`` column — an UPDATE row then leaves existing links
        # untouched instead of wiping them.
        req_ids, unknown = (None, []) if not has_req_column else ([], [])
        if has_req_column:
            for req_code in (data.get("requirements", "") or "").split(","):
                req_code = req_code.strip()
                if not req_code:
                    continue
                if req_code in req_by_code:
                    req_ids.append(req_by_code[req_code])
                else:
                    unknown.append(req_code)
        if unknown:
            errors.append("Unknown requirement code(s): " + ", ".join(unknown))

        clean = {
            "external_code": data.get("external_code", ""),
            "title": title,
            "description": data.get("description", ""),
            "actor": data.get("actor", ""),
            "priority": priority,
            "status": status,
            "remarks": data.get("remarks", ""),
            "requirement_ids": req_ids,
            "existing_id": existing_id,
        }
        preview.append(
            {"row": r_i, "data": clean, "errors": errors, "valid": not errors}
        )
    return preview


def import_usecases(poc, rows, user):
    """Create/update Use Cases from validated preview rows.

    A row whose ``code`` matched an existing use case of this POC (see
    ``parse_usecases_xlsx``) UPDATES it in place instead of creating a
    duplicate. Returns ``(created, updated)``.
    """
    created = updated = 0
    for row in rows:
        if not row.get("valid"):
            continue
        data = dict(row["data"])
        req_ids = data.pop("requirement_ids", None)
        existing_id = data.pop("existing_id", None)
        if existing_id:
            uc = UseCase.objects.get(pk=existing_id)
            for field_name, value in data.items():
                setattr(uc, field_name, value)
            uc.modified_by = user
            uc.save()
            updated += 1
            record_audit(uc, "usecase_updated", user, {})
        else:
            uc = UseCase.objects.create(
                poc=poc, created_by=user, modified_by=user, **data
            )
            created += 1
            record_audit(uc, "usecase_created", user, {"code": {"before": None, "after": uc.code}})
        if req_ids is not None:
            uc.requirements.set(req_ids)
    return created, updated


# ---------------------------------------------------------------------------
# Downloadable .xlsx templates (headers + example row + valid-options legend)
# ---------------------------------------------------------------------------
def _build_template_xlsx(headers, example_row, legend):
    """Build an .xlsx with a ``Template`` sheet and a ``Suggested options`` sheet.

    ``legend`` is a list of ``(column_name, [suggested_values])`` shown on the
    second sheet as a guide — these columns aren't restricted to the list.
    """
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    if example_row:
        ws.append(example_row)
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(14, width + 2)

    if legend:
        ws2 = wb.create_sheet("Suggested options")
        ws2.append(["Column", "Suggested values (not exclusive — other text is accepted)"])
        for cell in ws2[1]:
            cell.font = Font(bold=True)
        for column_name, values in legend:
            ws2.append([column_name, ", ".join(values)])
        for col in ws2.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws2.column_dimensions[col[0].column_letter].width = max(14, min(width + 2, 80))

    import io

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_requirements_template_xlsx(poc=None):
    """Return an in-memory .xlsx matching the Requirement import format.

    ``poc`` (optional) adds that POC's own previously-added custom
    classification values to the "Suggested options" sheet, alongside the
    built-in ones — the columns are not limited to either list (see module
    docstring).
    """
    headers = [
        "external_code",
        "sub_system",
        "req_gravity",
        "req_operation",
        "req_functional",
        "req_category",
        "description",
        "validation_criteria",
        "life_cycle_phase",
        "reference_documentations",
        "remarks",
        "use_cases",
    ]
    example_row = [
        "FR101",
        "PLC Controller",
        Requirement.Gravity.IMPOSES_MVP.label,
        Requirement.Operation.CONTROL_OPERATION.label,
        Requirement.Functional.PERFORMANCE.label,
        Requirement.Category.NORMAL_OPERATION.label,
        "The system shall display the real-time connector status (offline / available / plugged / "
        "charging / error) of every charging point in normal operation mode.",
        "The system shall respond within 200ms",
        Requirement.LifeCyclePhase.DRAFT.label,
        "SPEC-001",
        "",
        "",
    ]
    legend = [
        (field_name, [label for _, label in Requirement.field_choices(poc, field_name)])
        for field_name in (*_CLASSIFICATION_FIELDS, *_OPTIONAL_CLASSIFICATION_FIELDS)
    ]
    legend.append(
        ("use_cases", ["Comma-separated existing use case codes, e.g. POC-001-UC001, POC-001-UC002 — leave blank to link none"])
    )
    return _build_template_xlsx(headers, example_row, legend)


def build_usecases_template_xlsx():
    """Return an in-memory .xlsx matching the Use Case import format."""
    headers = [
        "external_code",
        "title",
        "description",
        "actor",
        "priority",
        "status",
        "remarks",
        "requirements",
    ]
    example_row = [
        "UC001",
        "Operator starts the process",
        "The operator triggers the process from the HMI",
        "Operator",
        UseCase.Priority.MEDIUM.label,
        UseCase.Status.DRAFT.label,
        "",
        "",
    ]
    legend = [
        ("priority", [label for _, label in UseCase.Priority.choices]),
        ("status", [label for _, label in UseCase.Status.choices]),
        ("requirements", ["Comma-separated existing requirement codes, e.g. REQ-001, REQ-002"]),
    ]
    return _build_template_xlsx(headers, example_row, legend)


# ---------------------------------------------------------------------------
# Full data export (round-trip: export → edit → re-import to update)
# ---------------------------------------------------------------------------
_REQUIREMENT_EXPORT_HEADERS = [
    "code",
    "external_code",
    "sub_system",
    "req_gravity",
    "req_operation",
    "req_functional",
    "req_category",
    "description",
    "validation_criteria",
    "life_cycle_phase",
    "reference_documentations",
    "remarks",
    "use_cases",
]

_USECASE_EXPORT_HEADERS = [
    "code", "external_code", "title", "description", "actor",
    "priority", "status", "remarks", "requirements",
]


def _requirement_export_rows(poc):
    """Headers matching the importer, plus a leading ``code`` column so a
    re-imported, edited row UPDATES the same requirement instead of creating
    a duplicate (see ``parse_requirements_xlsx``/``_md``). ``code`` is
    Light's own identifier — don't edit it; edit anything else, including
    ``external_code``."""
    rows = []
    for req in poc.requirements.prefetch_related("use_cases").order_by("code"):
        rows.append([
            req.code,
            req.external_code,
            req.sub_system,
            req.get_req_gravity_display(),
            req.get_req_operation_display(),
            req.get_req_functional_display(),
            req.get_req_category_display(),
            req.description,
            req.validation_criteria,
            req.get_life_cycle_phase_display(),
            req.reference_documentations,
            req.remarks,
            ", ".join(uc.code for uc in req.use_cases.all()),
        ])
    return rows


def _usecase_export_rows(poc):
    rows = []
    for uc in poc.use_cases.prefetch_related("requirements").order_by("code"):
        rows.append([
            uc.code,
            uc.external_code,
            uc.title,
            uc.description,
            uc.actor,
            uc.get_priority_display(),
            uc.get_status_display(),
            uc.remarks,
            ", ".join(r.code for r in uc.requirements.all()),
        ])
    return rows


def _build_export_xlsx(sheet_title, headers, rows):
    import io

    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(14, min(width + 2, 60))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_requirements_export_xlsx(poc):
    """Dump every current Requirement of ``poc`` as .xlsx — see ``build_requirements_export_md``
    for the Markdown equivalent (same columns, same round-trip semantics)."""
    return _build_export_xlsx("Requirements", _REQUIREMENT_EXPORT_HEADERS, _requirement_export_rows(poc))


def build_usecases_export_xlsx(poc):
    """Dump every current Use Case of ``poc`` as .xlsx — see ``build_requirements_export_xlsx``."""
    return _build_export_xlsx("Use Cases", _USECASE_EXPORT_HEADERS, _usecase_export_rows(poc))


def build_requirements_export_md(poc):
    """Dump every current Requirement of ``poc`` as a Markdown pipe table —
    same columns/semantics as ``build_requirements_export_xlsx``, readable and
    editable as plain text and reimportable via ``parse_requirements_md``."""
    return _write_md_table(_REQUIREMENT_EXPORT_HEADERS, _requirement_export_rows(poc))


def build_usecases_export_md(poc):
    """Dump every current Use Case of ``poc`` as a Markdown pipe table — see
    ``build_requirements_export_md``."""
    return _write_md_table(_USECASE_EXPORT_HEADERS, _usecase_export_rows(poc))
