"""
Typed blocks → python-docx Document → bytes (vendored).

Two modes:
  * template_path=None  → new document with built-in styles (Calibri).
  * template_path=...    → inherits styles/fonts/cover from the reference
    template (see template.py); adds PASS/FAIL cell colouring on top.
"""

import itertools
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from . import styles
from .template import prepare_from_template

# Cap embedded images so a large upload doesn't overflow the page width.
_MAX_IMAGE_WIDTH = Inches(6)

# A heading whose text starts with a code like "POC-119-UC001" or
# "POC-119-MM-AV-013" gets a bookmark on that token, so other sections (e.g. a
# Requirement's "Use cases:" line) can link straight to it.
_CODE_RE = re.compile(r"^([A-Za-z0-9]+(?:-[A-Za-z0-9]+){2,})")


def _anchor_name(code: str) -> str:
    """A Word-bookmark-safe name derived from a Requirement/UseCase code."""
    return re.sub(r"[^A-Za-z0-9_]", "_", code)


def build_docx(blocks: list[dict], metadata: dict | None = None,
               template_path=None, image_resolver=None) -> bytes:
    """Render typed blocks to a .docx.

    ``image_resolver`` (optional): ``url -> local file path | None``. When a
    block of type ``image`` is found its URL is resolved to a file and embedded;
    unresolved images fall back to their alt text.
    """
    use_template = template_path is not None
    if use_template:
        doc = prepare_from_template(template_path, metadata)
    else:
        doc = Document()
        _apply_base_styles(doc)
        if metadata:
            _add_cover(doc, metadata)

    table_style = _resolve_table_style(doc, use_template)
    bookmark_ids = itertools.count(1)

    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            _add_heading(doc, block, use_template, bookmark_ids)
        elif kind == "paragraph":
            _add_paragraph(doc, block)
        elif kind == "table":
            _add_table(doc, block, table_style, use_template)
        elif kind == "list":
            _add_list(doc, block)
        elif kind == "image":
            _add_image(doc, block, image_resolver)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _resolve_table_style(doc, use_template: bool) -> str:
    available = {s.name for s in doc.styles}
    if use_template and styles.TEMPLATE_TABLE_STYLE in available:
        return styles.TEMPLATE_TABLE_STYLE
    return "Table Grid" if "Table Grid" in available else "Normal Table"


def _apply_base_styles(doc):
    font = doc.styles["Normal"].font
    font.name = styles.FONT_NAME
    font.size = styles.FONT_SIZE_BODY


def _add_cover(doc, meta: dict):
    title = meta.get("title", "")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(title))
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = styles.COLOR_ACCENT

    for key, val in meta.items():
        if key == "title" or val in (None, ""):
            continue
        row = doc.add_paragraph()
        label = row.add_run(f"{key.capitalize()}: ")
        label.bold = True
        row.add_run(str(val))

    doc.add_page_break()


def _add_heading(doc, block: dict, use_template: bool, bookmark_ids):
    level = min(int(block["level"]), 4)
    p = doc.add_heading(block["text"], level=level)
    if not use_template and p.runs:
        run = p.runs[0]
        run.font.color.rgb = styles.COLOR_ACCENT
        run.font.size = styles.SIZE_BY_LEVEL.get(level, styles.SIZE_BY_LEVEL[3])
    match = _CODE_RE.match(block["text"])
    if match:
        _add_bookmark(p, _anchor_name(match.group(1)), bookmark_ids)


def _add_bookmark(paragraph, name, bookmark_ids):
    bid = str(next(bookmark_ids))
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_paragraph(doc, block: dict):
    p = doc.add_paragraph()
    _render_runs(p, block.get("children", []))


# Word has no native SVG support in the add_picture path, so an SVG is
# rasterized to PNG at generation time (never in the editor, which just shows
# it via a plain <img> — browsers render SVG natively). The render is capped
# to a fixed pixel width rather than a zoom multiplier, so a pathological
# SVG that declares a huge canvas can't blow up memory/time.
_SVG_RASTER_WIDTH_PX = 1600


def _svg_to_png_bytes(svg_path: str) -> bytes:
    import resvg_py

    png = resvg_py.svg_to_bytes(
        svg_path=svg_path, width=_SVG_RASTER_WIDTH_PX, background="#FFFFFF"
    )
    return bytes(png)


def _add_image(doc, block: dict, image_resolver):
    """Embed an image (resolved to a local path) or fall back to its alt text."""
    url = block.get("url", "")
    path = image_resolver(url) if (image_resolver and url) else None
    if path:
        try:
            if path.lower().endswith(".svg"):
                picture = doc.add_picture(BytesIO(_svg_to_png_bytes(path)))
            else:
                picture = doc.add_picture(path)
            # Downscale only if it overflows the page; never upscale small images.
            if picture.width > _MAX_IMAGE_WIDTH:
                ratio = _MAX_IMAGE_WIDTH / picture.width
                picture.width = _MAX_IMAGE_WIDTH
                picture.height = int(picture.height * ratio)
            return
        except Exception:  # noqa: BLE001 — unreadable/unconvertible image → alt-text fallback
            pass
    alt = block.get("alt") or url
    if alt:
        p = doc.add_paragraph()
        run = p.add_run(f"[image: {alt}]")
        run.italic = True


def _add_list(doc, block: dict):
    available = {s.name for s in doc.styles}
    if block.get("ordered"):
        style = "List Number" if "List Number" in available else "List Paragraph"
    else:
        style = "List Bullet" if "List Bullet" in available else "List Paragraph"
    for item_runs in block["items"]:
        p = doc.add_paragraph(style=style)
        _render_runs(p, item_runs)


def _render_runs(paragraph, runs: list[dict]):
    for rd in runs:
        link = rd.get("link")
        if link and link.startswith("#"):
            _add_internal_hyperlink(
                paragraph, rd.get("text", ""), _anchor_name(link[1:]), rd.get("bold", False)
            )
            continue
        run = paragraph.add_run(rd.get("text", ""))
        run.bold = rd.get("bold", False)
        run.italic = rd.get("italic", False)
        if run.bold:
            run.font.color.rgb = styles.COLOR_ACCENT
        if rd.get("code"):
            run.font.name = "Consolas"
        if link:
            run.font.underline = True
            run.font.color.rgb = styles.COLOR_ACCENT


def _add_internal_hyperlink(paragraph, text, anchor, bold=False):
    """A clickable in-document link (Word bookmark reference), e.g. a
    Requirement's "Use cases:" entry jumping to that Use Case's heading."""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)

    run_el = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "{:02X}{:02X}{:02X}".format(*styles.COLOR_ACCENT))
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    if bold:
        rpr.append(OxmlElement("w:b"))
    run_el.append(rpr)

    text_el = OxmlElement("w:t")
    text_el.text = text
    text_el.set(qn("xml:space"), "preserve")
    run_el.append(text_el)

    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def _add_table(doc, block: dict, table_style: str, use_template: bool):
    headers = block["headers"]
    rows = block["rows"]
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = table_style
    _enable_first_row_format(table)

    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        para = cell.paragraphs[0]
        run = para.runs[0] if para.runs else para.add_run("")
        run.bold = True
        if not use_template:
            run.font.color.rgb = styles.COLOR_HEADER_FG
            _set_cell_bg(cell, styles.COLOR_HEADER_BG)

    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = text
            upper = text.strip().upper()
            if upper.startswith("PASS") or upper in ("OK", "✓"):
                _set_cell_bg(cell, styles.COLOR_PASS_BG)
            elif (
                upper.startswith("FAIL")
                or upper.startswith("NOT PASSED")
                or upper in ("KO", "NOK", "✗")
            ):
                _set_cell_bg(cell, styles.COLOR_FAIL_BG)


def _enable_first_row_format(table):
    tblPr = table._tbl.tblPr
    look = tblPr.find(qn("w:tblLook"))
    if look is None:
        look = OxmlElement("w:tblLook")
        tblPr.append(look)
    look.set(qn("w:firstRow"), "1")
    look.set(qn("w:lastRow"), "0")
    look.set(qn("w:firstColumn"), "0")
    look.set(qn("w:lastColumn"), "0")
    look.set(qn("w:noHBand"), "0")
    look.set(qn("w:noVBand"), "1")


def _set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
