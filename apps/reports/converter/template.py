"""
Reference-template support (vendored).

Opens a reference .docx (inheriting its styles, fonts, cover page, headers/
footers, TOC and section setup), strips the sample content and leaves the
document ready for the builder to append the converted Markdown body.
"""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from . import styles


def list_templates() -> list[dict]:
    if not styles.TEMPLATES_DIR.is_dir():
        return []
    items = []
    for path in sorted(styles.TEMPLATES_DIR.glob("*.docx")):
        if path.name.startswith("~$"):
            continue
        items.append({"id": path.name, "name": path.stem})
    return items


def resolve_template(name):
    if not name or str(name).lower() == "none":
        return None
    candidate = styles.TEMPLATES_DIR / name
    if candidate.exists():
        return candidate
    p = Path(name)
    return p if p.exists() else None


def prepare_from_template(template_path, metadata: dict | None = None) -> Document:
    """Document based on the template, with sample content removed.

    Preserves everything before the first marker block (default ``Heading 1``):
    cover page, document-control page and TOC.
    """
    doc = Document(str(template_path))
    _strip_sample_content(doc)
    if metadata:
        _fill_cover_fields(doc, metadata)
    _update_fields_on_open(doc)
    return doc


def _strip_sample_content(doc: Document):
    body = doc.element.body
    marker = styles.TEMPLATE_CONTENT_MARKER_STYLE

    start = None
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            if para.style and para.style.name == marker:
                start = child
                break

    if start is None:
        return

    to_remove = []
    found = False
    for child in body.iterchildren():
        if child is start:
            found = True
        if not found:
            continue
        if child.tag in (qn("w:p"), qn("w:tbl")):
            to_remove.append(child)
    for el in to_remove:
        body.remove(el)


def _fill_cover_fields(doc: Document, metadata: dict):
    paragraphs = doc.paragraphs
    for i, para in enumerate(paragraphs):
        label = para.text.strip()
        key = styles.COVER_FIELD_LABELS.get(label)
        if not key or key not in metadata:
            continue
        value = str(metadata[key])
        for nxt in paragraphs[i + 1:]:
            if nxt.text.strip():
                _set_paragraph_text(nxt, value)
                break


def _set_paragraph_text(paragraph: Paragraph, text: str):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for extra in runs[1:]:
        extra.text = ""


def _update_fields_on_open(doc: Document):
    """Mark the document so Word updates fields (TOC) on open."""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)
